"""Lê a planilha de projeções da consultoria (ex.: Tendências) para o Radar.

Formato esperado (primeira aba): primeira linha = "Índices" + anos (2015, 2016, ...); cada linha
seguinte = um indicador, com histórico e projeção na mesma linha. Não há marcação de quais anos
são projeção: por padrão, a projeção começa no ano da edição (ex.: arquivo "102026_Dexco.xlsx" =
edição out/2026, projeção a partir de 2026), ou no ano passado em --proj-desde.

Cada indicador vai para um tópico do Radar conforme projecoes_map.json (por trecho do nome).
Indicadores com "%" no nome são taxas em nível; os demais são variação % a/a.
"""
import json, os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
MES = ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']

def edicao_do_nome(path):
    m = re.search(r'(?<!\d)(0[1-9]|1[0-2])(20\d\d)(?!\d)', os.path.basename(path))
    return (int(m.group(1)), int(m.group(2))) if m else None

def ler(path, desde=None, fonte='Tendências'):
    mapa = json.load(open(os.path.join(HERE, 'projecoes_map.json'), encoding='utf-8'))
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    anos, cols = [], []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is not None and re.fullmatch(r'\d{4}', str(v).strip()):
            anos.append(int(str(v).strip())); cols.append(c)
    if not anos:
        raise SystemExit(f'Projeções: a primeira linha de "{ws.title}" não tem os anos nas colunas.')
    ed = edicao_do_nome(path)
    if desde is None:
        desde = ed[1] if ed else max(a for a in anos if a <= __import__('datetime').date.today().year)
    itens, sem_topico = [], []
    for r in range(2, ws.max_row + 1):
        nome = ws.cell(r, 1).value
        if not nome or not str(nome).strip():
            continue
        nome = str(nome).strip()
        regra = next((m for m in mapa if re.search(m['padrao'], nome, re.I)), None)
        nivel = '%' in nome if not regra or 'tipo' not in regra else regra['tipo'] == 'nivel'
        dados = []
        for c in cols:
            cel = ws.cell(r, c)
            v = cel.value
            if isinstance(v, (int, float)):
                # células sem formato de % trazem o número já em % (ex.: 5,7 = 5,7%)
                v = v if '%' in (cel.number_format or '') else v / 100
            else:
                v = None
            dados.append(v)
        item = {'id': f'tend{r}', 'nome': nome, 'titulo': regra['titulo'] if regra else nome,
                'sec': regra['sec'] if regra else None, 'tipo': 'nivel' if nivel else 'var',
                'labels': [str(a) for a in anos], 'data': dados, 'cel': f'{ws.title}!A{r}'}
        itens.append(item)
        if not regra:
            sem_topico.append(nome)
    return {'fonte': fonte, 'edicao': f'{MES[ed[0]-1]}/{str(ed[1])[2:]}' if ed else None,
            'arquivo': os.path.basename(path), 'desde': desde, 'itens': itens}, sem_topico

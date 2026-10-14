import openpyxl,warnings,json,re,sys
warnings.filterwarnings('ignore')
from openpyxl.utils import get_column_letter as L
import os
wb=openpyxl.load_workbook(os.environ['PEERS_XLSX'],read_only=True,data_only=True)
def isyear(v):
    if isinstance(v,(int,float)) and 2000<=v<=2035 and float(v).is_integer(): return True
    if isinstance(v,str) and re.fullmatch(r"\s*(20\d\d|\d?T\d\d.*LTM.*|.*LTM.*|20\d\dE)\s*",v): return True
    return False
out={}
for name in sys.argv[1:]:
    ws=wb[name]; rows=[list(r) for r in ws.iter_rows(values_only=True)]
    blocks=[]; i=0
    while i<len(rows):
        r=rows[i]+[None]*40
        found=None
        for c in range(0,30):
            if isinstance(r[c],str) and r[c].strip() and not isyear(r[c]) and isyear(r[c+1]) and isyear(r[c+2]):
                found=c;break
        if found is None: i+=1;continue
        c=found; yrs=[]; k=c+1
        while k<len(r) and isyear(r[k]): yrs.append(str(r[k]).strip() if not isinstance(r[k],(int,float)) else str(int(r[k]))); k+=1
        title=r[c].strip(); comp=[]; j=i+1
        while j<len(rows):
            rr=rows[j]+[None]*40
            nm=rr[c]
            if nm is None or (isinstance(nm,str) and not nm.strip()): break
            if any(isinstance(rr[x],str) and isyear(rr[x]) for x in range(c+1,c+3)): break
            vals=[]
            for x in range(c+1,c+1+len(yrs)):
                v=rr[x]; vals.append(float(v) if isinstance(v,(int,float)) and not isinstance(v,bool) else None)
            comp.append(dict(name=str(nm).strip(),country=(str(rr[c-1]).strip() if c>0 and rr[c-1] and len(str(rr[c-1]))<=4 else None),vals=vals,row=j+1))
            j+=1
        blocks.append(dict(title=title,row=i+1,col=L(c+1),years=yrs,comp=comp))
        i=j
    out[name]=blocks
json.dump(out,open('blocks.json','w'),ensure_ascii=False)
for n,bs in out.items():
    print('=====',n,len(bs))
    for b in bs: print(f" r{b['row']} {b['col']} {b['title'][:40]!r} yrs={b['years'][0]}..{b['years'][-1]}({len(b['years'])}) n={len(b['comp'])} [{', '.join(x['name'][:14] for x in b['comp'][:11])}]")

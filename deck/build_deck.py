import json,copy,re,statistics
from pptx import Presentation
from pptx.util import Inches,Pt,Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_MARKER_STYLE,XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
import sys
# uso: python build_deck.py <peers_deck.json> <template.pptx> <saida.pptx>
D=json.load(open(sys.argv[1]))
prs=Presentation(sys.argv[2])
S=list(prs.slides)
cover,menu,s_mad,s_lou,s_met,s_rev,closing=S
REF='2025'; Y0=2019
RGB=lambda h:RGBColor.from_string(h)
SEGCOL={'loucas':'203864','metais':'2F5597','madeira':'548235','revest':'ED7D31'}
COL={
 'loucas':{'icasa':'70AD47','roca_do_brasil':'0070C0','astra':'FF0000','toto':'FFC000','lixil':'ED7D31','roca':'A5A5A5'},
 'metais':{'docol':'4BACC6','fabrimar':'FF3300','masco':'EA9618','gessi':'FFC000','geberit':'4472C4','villeroy_boch':'70AD47'},
 'madeira':{'berneck':'4BACC6','eucatex':'FF3300','arauco_br':'9BBB59','guararapes_paineis':'8064A2','placas_do_brasil':'D7428B','repinho':'747474','sudati':'FFC000','floraplac':'764B0B','fibraplac':'AF2468','dexco_paineis_verticalizados':'404040'},
 'revest':{'portobello':'0070C0','villagres':'70AD47','biancogres':'4BACC6','embramaco':'ED7D31','mohawk':'FF0000','porcelanosa':'8064A2','grupo_lamosa':'FFC000'},
}
EXTRA={'grupo roca':'0070C0','docol':'4BACC6','florestal':'548235','painéis':'2F5597','colômbia':'ED7D31','madeira br':'7F7F7F','egger':'2F5597','west fraser timber':'ED7D31','masisa':'70AD47','fiplasto':'7F7F7F','deca metais':'000000','fabrimar':'FF3300'}
AX={
 'loucas':{'rol':(None,80000,1),'ebitda':(None,8000,1),'ll':(-700,3000,1),'ci':(0,60000,1),'me':(-.3,.7),'roic':(None,.5),'giro':(None,7),'alav':(None,7),'mb':(0,1),'ml':(None,.5),'cpv':(0,1),'dga':(0,.6),'dv':(-.1,.5),'pmp':(None,250),'wk':(None,.5),'pmr':(None,100),'pme':(None,250),'fcl':(None,3),'varwk':(-2,2),'capexe':(-2,2),'outras':(-2,2),'capexr':(0,.15),
   'Market share – Competitivo':(0,None),'Market share – Médio/Luxo':(0,None),'Preço médio – Competitivo':(100,300),'Preço médio – Médio/Luxo':(500,1200)},
 'metais':{'rol':(0,50000,1),'ebitda':(0,10000,1),'ll':(0,10000,1),'ci':(0,20000,1),'me':(0,1),'roic':(0,.6),'alav':(-5,5),'giro':(0,15),'mb':(0,1),'ml':(-.3,.3),'cpv':(0,1),'dga':(0,.5),'dv':(0,.5),'fcl':(-1.5,1.5),'varwk':(-2,2),'capexe':(-2.5,.5),'outras':(-1.3,1.3),'pmp':(0,400),'wk':(-.1,.8),'pmr':(None,100),'pme':(0,400),'capexr':(-.02,.15),
   'Preço médio – Competitivo':(100,None),'Preço médio – Médio/Luxo':(300,None)},
 'madeira':{'rol':(0,7000),'roic':(0,.25),'me':(0,.4),'giro':(.2,1.8),'wk':(0,.4),'pmr':(0,150),'pme':(0,150),'pmp':(0,150),'ROIC: players internacionais vs Dexco':(None,.25)},
 'revest':{'rol':(0,70000,1),'ebitda':(0,12000,1),'ll':(-1000,12000,1),'ci':(0,70000,1),'alav':(-2,6),'me':(-.1,.9),'roic':(-.1,.6),'giro':(None,3),'dv':(-.1,.5),'dga':(0,.5),'cpv':(0,1),'wk':(0,1),'pmr':(None,200),'pme':(None,200),'pmp':(None,200),'fcl':(-1,2.2),'mb':(0,1),'ml':(-.4,.4),'capexr':(-.1,.4)},
}
BETTER={'rol':1,'ebitda':1,'ll':1,'roic':1,'me':1,'mb':1,'ml':1,'giro':1,'alav':-1,'cpv':-1,'dga':-1,'dv':-1,'wk':-1,'pmr':-1,'pme':-1,'pmp':1,'fcl':1}
def br(x,d=1):
    s=f'{abs(x):,.{d}f}'.replace(',','X').replace('.',',').replace('X','.');return s
def fmt(v,f,paren=True):
    if v is None: return '–'
    neg=v<0
    if f=='mi': s=br(v,1)
    elif f=='pct': s=br(v*100,1)+'%'
    elif f=='x': s=br(v,1)+'x'
    elif f=='d': s=br(v,0)
    elif f=='brl': s='R$ '+br(v,0)
    else: s=br(v,1)
    if neg: s=('('+s.replace('%','').replace('x','')+')'+('%' if f=='pct' else 'x' if f=='x' else '')) if paren else '−'+s
    return s
def valAt(m,k,ano=REF,ns=False):
    v=m['series'].get(k);best=(None,None,None)
    if not v: return best if ns else best[:2]
    ex=(m.get('excl') or {}).get(k,{})
    for y,x in zip(m['years'],v):
        if 'LTM' in y or int(y)>int(ano): continue
        if x is not None: best=(x,y,None)
        elif y in ex: best=(None,y,ex[y])
    return best if ns else best[:2]
def rank(seg,key):
    s=D[seg];m=s['metrics'].get(key)
    if not m: return None
    vals=[(k,)+valAt(m,k) for k in s['ents'] if not k.startswith('dexco_')]
    vals=[x for x in vals if x[1] is not None and int(REF)-int(x[2])<=1]
    dex=next((x for x in vals if x[0]=='dexco'),None);peers=[x[1] for x in vals if x[0]!='dexco']
    med=statistics.median(peers) if peers else None;pos=None
    if dex and BETTER.get(key):
        srt=sorted(vals,key=lambda x:-BETTER[key]*x[1]);pos=[x[0] for x in srt].index('dexco')+1
    return dict(dex=dex,med=med,pos=pos,n=len(vals))
def tb(slide,x,y,w,h,text,size=11,bold=False,color='262626',align=PP_ALIGN.LEFT,font='Roboto',anchor=MSO_ANCHOR.MIDDLE,italic=False):
    t=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=t.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.04);tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=anchor
    p=tf.paragraphs[0];p.alignment=align;r=p.add_run();r.text=text;f=r.font;f.size=Pt(size);f.bold=bold;f.italic=italic;f.name=font;f.color.rgb=RGB(color)
    return t
# ---------- matrizes-resumo ----------
SUMROWS=['rol','ebitda','me','fcl','giro','ci','alav','roic']
def fill_summary(slide,seg,cols,rows,colw):
    s=D[seg]
    for k,x in cols:
        ys=[valAt(s['metrics'][m],k)[1] for m in SUMROWS if m in s['metrics'] and valAt(s['metrics'][m],k)[1]]
        yy=max(ys) if ys else '–'
        tb(slide,x,1.5,colw,.22,yy,9,False,'7F7F7F',PP_ALIGN.CENTER,italic=True)
        for m,(y,h) in zip(SUMROWS,rows):
            M=s['metrics'].get(m)
            if not M: tb(slide,x,y,colw,h,'–',11,False,'7F7F7F',PP_ALIGN.CENTER);continue
            v,vy,nsx=valAt(M,k,ns=True)
            t=tb(slide,x,y,colw,h,'n.s.' if nsx else fmt(v,M['fmt']),11,k=='dexco','262626' if v is not None else 'A5A5A5',PP_ALIGN.CENTER)
            if v is not None and vy!=REF:
                r=t.text_frame.paragraphs[0].add_run();r.text=' '+vy;r.font.size=Pt(7.5);r.font.color.rgb=RGB('8C8C8C');r.font.name='Roboto';r.font.italic=True
    # posição Dexco
    xr=cols[-1][1]+colw+0.02
def title(slide,text,sub=None):
    tb(slide,1.05,0.2,12,.45,text.upper(),18,True,'000000',font='Roboto Black')
    if sub: tb(slide,1.05,0.6,4.5,.4,sub,9,False,'595959',anchor=MSO_ANCHOR.TOP)
def footnote(slide,text):
    tb(slide,1.05,7.13,12.05,.3,text,7.5,False,'7F7F7F',anchor=MSO_ANCHOR.TOP)
def find(slide,name): return next(sh for sh in slide.shapes if sh.name==name)
def settext(sh,text):
    p=sh.text_frame.paragraphs[0];rs=p.runs
    rs[0].text=text
    for r in rs[1:]: r.text=''

LBL_A=[(2.07,.39),(2.61,.39),(3.15,.39),(3.67,.39),(4.22,.39),(4.75,.39),(5.30,.39),(5.84,.39)] # madeira
LBL_B=[(2.08,.39),(2.73,.39),(3.37,.39),(4.01,.39),(4.66,.39),(5.30,.39),(5.94,.39),(6.59,.39)]
# Madeira
settext(find(s_mad,'Rectangle 8'),'FCL/EBITDA')
mad_cols=[('dexco',2.51),('berneck',3.57),('eucatex',4.63),('arauco_br',5.69),('guararapes_paineis',6.74),('placas_do_brasil',7.81),('repinho',8.87),('sudati',9.93),('floraplac',10.97),('fibraplac',12.02)]
fill_summary(s_mad,'madeira',mad_cols,LBL_A,.98)
lou_cols=[('dexco',2.67),('icasa',4.41),('roca_do_brasil',6.15),('astra',7.89),('toto',9.63),('lixil',11.37)]
fill_summary(s_lou,'loucas',lou_cols,LBL_B,1.59)
met_cols=[('dexco',2.52),('docol',4.06),('fabrimar',5.6),('masco',7.14),('gessi',8.67),('geberit',10.21),('villeroy_boch',11.75)]
fill_summary(s_met,'metais',met_cols,LBL_B,1.32)
# Revestimentos: monta a matriz copiando a estrutura de Metais
for sh in s_met.shapes:
    if sh.shape_type==13: continue
    if sh.name in ('Rectangle 4','TextBox 14','Rectangle 16') or (sh.name.startswith('Rectangle 2') and sh.has_text_frame and sh.text_frame.text.strip()):
        s_rev.shapes._spTree.append(copy.deepcopy(sh._element))
# (os rótulos das linhas foram copiados; agora cabeçalhos coloridos como em Madeira)
tag=find(s_mad,'Rectangle 6')
rev_ents=[('dexco','DX Revest.','404040'),('portobello','Portobello','0070C0'),('villagres','Villagres','70AD47'),('biancogres','Biancogres','4BACC6'),('embramaco','Embramaco','ED7D31'),('mohawk','Mohawk','FF0000'),('porcelanosa','Porcelanosa','8064A2'),('grupo_lamosa','Grupo Lamosa','FFC000')]
step=(13.07-2.52)/8;w=step-.1;rev_cols=[]
for i,(k,n,c) in enumerate(rev_ents):
    el=copy.deepcopy(tag._element);s_rev.shapes._spTree.append(el);sh=s_rev.shapes[-1]
    x=2.52+i*step;sh.left=Inches(x);sh.top=Inches(1.07);sh.width=Inches(w);sh.height=Inches(.39)
    sh.fill.solid();sh.fill.fore_color.rgb=RGB(c);settext(sh,n)
    rev_cols.append((k,x))
fr=next(sh for sh in s_rev.shapes if sh.name=='Rectangle 16');fr.left=Inches(2.52-.06);fr.width=Inches(w+.12)
fill_summary(s_rev,'revest',rev_cols,LBL_B,w)
SUMNOTE='Valores de 2025 em R$ milhões; o ano aparece ao lado quando o dado mais recente é anterior. Giro de ativos = ROL / capital investido. Endividamento = dívida líquida / EBITDA. Negativos entre parênteses. n.s. = valor excluído por erro ou outlier (ver Ajustes na base).'
for sl,seg in [(s_lou,'loucas'),(s_met,'metais'),(s_mad,'madeira'),(s_rev,'revest')]:
    title(sl,f"Peers {D[seg]['name']} – Resumo {REF}")
    footnote(sl,SUMNOTE+(' EBITDA de Madeira = ROL × margem EBITDA (a planilha não traz o valor absoluto).' if seg=='madeira' else '')+(' Endividamento da Dexco = alavancagem consolidada.'))

# ---------- gráficos ----------
NUMF={'mi':'#,##0','pct':'0%','x':'0.0"x"','d':'0','brl':'"R$" #,##0'}
def ent_color(seg,k):
    if k=='dexco': return '000000'
    return COL[seg].get(k,'A5A5A5')
def add_chart(slide,x,y,w,h,labels,series,fmt_,ax,bar):
    cd=CategoryChartData();cd.categories=[l.replace(' LTM',' LTM') for l in labels]
    for s in series: cd.add_series(s['name'],[None if v is None else float(v) for v in s['data']])
    ct=XL_CHART_TYPE.COLUMN_CLUSTERED if bar else XL_CHART_TYPE.LINE_MARKERS
    gf=slide.shapes.add_chart(ct,Inches(x),Inches(y),Inches(w),Inches(h),cd);ch=gf.chart
    ch.has_title=False;ch.font.size=Pt(8);ch.font.name='Roboto';ch.font.color.rgb=RGB('595959')
    ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False;ch.legend.font.size=Pt(7.5)
    va=ch.value_axis;va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=RGB('E7E6E6');va.major_gridlines.format.line.width=Pt(.5)
    va.format.line.fill.background();va.tick_labels.number_format=NUMF.get(fmt_,'General');va.tick_labels.number_format_is_linked=False;va.tick_labels.font.size=Pt(8)
    if ax:
        if ax[0] is not None: va.minimum_scale=ax[0]
        if ax[1] is not None: va.maximum_scale=ax[1]
    ca=ch.category_axis;ca.tick_labels.font.size=Pt(8);ca.format.line.color.rgb=RGB('BFBFBF');ca.tick_label_position=XL_TICK_LABEL_POSITION.LOW;ca.has_major_gridlines=False
    plot=ch.plots[0]
    if bar: plot.gap_width=60;plot.overlap=0
    for s,ps in zip(series,plot.series):
        c=RGB(s['color'])
        if bar:
            ps.format.fill.solid();ps.format.fill.fore_color.rgb=c;ps.invert_if_negative=False
        else:
            ps.smooth=False;ln=ps.format.line;ln.color.rgb=c;ln.width=Pt(s.get('w',1.5))
            if s.get('dash'): ln.dash_style=s['dash']
            if s.get('dex'):
                ps.marker.style=XL_MARKER_STYLE.CIRCLE;ps.marker.size=5;ps.marker.format.fill.solid();ps.marker.format.fill.fore_color.rgb=c;ps.marker.format.line.color.rgb=c
            else: ps.marker.style=XL_MARKER_STYLE.NONE
    return ch
def metric_series(seg,key):
    s=D[seg];m=s['metrics'][key];idx=[i for i,y in enumerate(m['years']) if (int(y[:4]) if 'LTM' not in y else 2026)>=Y0]
    labels=[m['years'][i] for i in idx];series=[]
    order=['dexco']+[k for k in s['ents'] if k!='dexco']
    for k in order:
        if k not in m['series']: continue
        data=[m['series'][k][i] for i in idx]
        if all(v is None for v in data): continue
        e=s['ents'][k];dex=k=='dexco'
        series.append(dict(name=e['name'],data=data,color=ent_color(seg,k),dex=dex,w=2.75 if dex else 1.5,dash=MSO_LINE_DASH_STYLE.DASH if k.startswith('dexco_') else None))
    return labels,series
import math
def focus_ax(series,bar):
    dv=[v for s in series if s.get('dex') for v in s['data'] if v is not None]
    if not dv: return None
    lo,hi=min(dv),max(dv)
    if bar or (lo>0 and lo<hi*.35): lo=min(lo,0)
    hi=max(hi,0)
    pad=max((hi-lo)*.45,abs(hi or lo)*.25,1e-6)
    mn=lo-(pad if lo<0 else 0);mx=hi+pad
    r=(mx-mn)/5;p=10**math.floor(math.log10(r));m=r/p;st=(1 if m<=1 else 2 if m<=2 else 2.5 if m<=2.5 else 5 if m<=5 else 10)*p
    mn=math.floor(mn/st)*st;mx=math.ceil(mx/st)*st
    if lo>=0 and mn<0: mn=0
    return (mn,mx)
def data_ax(series,bar):
    v=[x for s in series for x in s['data'] if x is not None]
    if not v: return None
    lo,hi=min(v),max(v)
    if bar: lo=min(lo,0);hi=max(hi,0)
    if hi==lo: hi+=abs(hi)*.1 or 1;lo-=abs(lo)*.1 or 1
    r=(hi-lo)/5;p=10**math.floor(math.log10(r));m=r/p;st=(1 if m<=1 else 2 if m<=2 else 2.5 if m<=2.5 else 5 if m<=5 else 10)*p
    mn=math.floor(lo/st)*st;mx=math.ceil(hi/st)*st
    if not bar and lo>=0 and mn<0: mn=0
    return (mn,mx)
def clipped_note(series,labels,ax,f):
    if not ax: return []
    out=[]
    for s in series:
        hits=[(l,v) for l,v in zip(labels,s['data']) if v is not None and ((ax[0] is not None and v<ax[0]) or (ax[1] is not None and v>ax[1]))]
        if not hits: continue
        out.append(f"{s['name']} ({len(hits)} anos)" if len(hits)>2 else s['name']+' '+', '.join(f"{l}: {fmt(v,f,False)}" for l,v in hits))
    return out
GROUPS=[('escala','Escala',['rol','ebitda','ll','ci'],'Receita, resultado e capital investido (R$ milhões)'),
 ('foco','Escala – Foco Dexco',['rol','ebitda','ll','ci'],'Mesmos gráficos com o eixo ajustado aos valores da Dexco (R$ milhões)'),
 ('rent','Rentabilidade',['roic','me','giro','alav'],'Retorno, margem, giro e alavancagem'),
 ('custos','Margens e custos',['mb','ml','cpv','dga','dv'],'Margens e estrutura de custos sobre a receita'),
 ('giro','Capital de giro',['wk','pmr','pme','pmp'],'Capital de giro / ROL e prazos médios (dias)'),
 ('caixa','Fluxo de caixa',['fcl','varwk','capexe','outras','capexr'],'Conversão de EBITDA em caixa'),
 ('setor','Setorial',None,'Mercado, participação e preço')]
NAV=[('menu','Menu'),('resumo','Resumo'),('escala','Escala'),('foco','Foco Dexco'),('rent','Rentab.'),('custos','Custos'),('giro','Giro'),('caixa','Caixa'),('setor','Setorial'),('ajustes','Ajustes')]
layout=s_lou.slide_layout
SEGSLIDES={};NAVSHAPES=[]
summary={'loucas':s_lou,'metais':s_met,'madeira':s_mad,'revest':s_rev}
def new_slide():
    sl=prs.slides.add_slide(layout)
    for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
    return sl
def grid(n):
    if n<=4: cols,rows=2,2
    else: cols,rows=3,2
    W=(13.1-1.05-(cols-1)*.25)/cols;H=(7.0-1.05-(rows-1)*.2)/rows
    return [(1.05+(i%cols)*(W+.25),1.05+(i//cols)*(H+.2),W,H) for i in range(n)]
for seg in ['loucas','metais','madeira','revest']:
    s=D[seg];SEGSLIDES[seg]={'resumo':summary[seg]}
    for gid,gname,keys,sub in GROUPS:
        if gid=='setor':
            items=[('S',b) for b in s['sector']]
        else:
            items=[('M',k) for k in keys if k in s['metrics']]
        if not items: continue
        sl=new_slide();SEGSLIDES[seg][gid]=sl
        title(sl,f"Peers {s['name']} – {gname}",sub)
        notes=[];refs=[];exnotes=[]
        for (kind,it),(x,y,w,h) in zip(items,grid(len(items))):
            if kind=='M':
                m=s['metrics'][it];labels,series=metric_series(seg,it);ax=None;bar=it in ('rol','ebitda','ll','ci') and seg!='madeira'
                if gid=='foco': ax=focus_ax(series,bar)
                ttl=m['label']+(' ¹' if m.get('calc') else '');f=m['fmt']
                r=rank(seg,it)
                dns=valAt(m,'dexco',ns=True)
                if dns[2]:
                    rd=f"{s['ents']['dexco']['name']} {dns[1]}: n.s. (valor de {fmt(dns[2]['v'],f,False)} excluído como {dns[2]['kind']}) · mediana peers: {fmt(r['med'],f,False) if r else '–'}"
                elif r and r['dex']:
                    rd=f"{s['ents']['dexco']['name']} {r['dex'][2]}: {fmt(r['dex'][1],f,False)} · mediana peers: {fmt(r['med'],f,False)}"
                else: rd=f"{s['ents']['dexco']['name']}: sem dado"
                refs.append(m['ref'])
            else:
                b=it;yrs=b['years'];idx=[i for i,yv in enumerate(yrs) if int(yv[:4])>=Y0];labels=[yrs[i] for i in idx];series=[];f=b['fmt'];ax=AX[seg].get(b['title']);bar=False
                occ='ocupa' in b['title'].lower()
                for x_ in b['series']:
                    nm=x_['name'];data=[x_['vals'][i] for i in idx];low=nm.lower()
                    if occ:
                        c='ED7D31' if 'MDP' in nm else '2F5597';series.append(dict(name=nm,data=data,color=c,w=2.25 if not x_['ref'] else 1.5,dash=MSO_LINE_DASH_STYLE.DASH if x_['ref'] else None));continue
                    if x_['ref']: series.append(dict(name=nm,data=data,color='7F7F7F',w=1.5,dash=MSO_LINE_DASH_STYLE.DASH));continue
                    if x_['dex']: series.append(dict(name=nm,data=data,color='000000',w=2.75,dex=True));continue
                    if low.startswith('outros'): series.append(dict(name=nm,data=data,color='A5A5A5',w=1.25,dash=MSO_LINE_DASH_STYLE.ROUND_DOT));continue
                    key=next((k for k,e in s['ents'].items() if low.startswith(e['name'].lower()) or e['name'].lower().startswith(low)),None)
                    c=COL[seg].get(key) if key else None
                    c=c or EXTRA.get(low.strip()) or 'A5A5A5'
                    series.append(dict(name=nm,data=data,color=c,w=1.75))
                ttl=b['title'];rd=b['unit'];refs.append(b['ref'])
            if gid!='foco': ax=data_ax(series,bar)
            tb(sl,x,y,w,.28,ttl,11.5,True,'000000',font='Roboto Black')
            pass
            add_chart(sl,x,y+.3,w,h-.3,labels,series,f,ax,bar)
            cl=clipped_note(series,labels,ax,f)
            if kind=='M':
                for ek,ys in (m.get('excl') or {}).items():
                    it=[(y_,o) for y_,o in ys.items() if (int(y_[:4]) if 'LTM' not in y_ else 2026)>=Y0]
                    if it: exnotes.append(f"{ttl} – {s['ents'][ek]['name']} "+(f"{len(it)} anos" if len(it)>2 else ', '.join(f"{y_}: {fmt(o['v'],f,False)}" for y_,o in it))+f" ({'/'.join(sorted(set(o['kind'] for _,o in it)))})")
            if cl: notes.append(f"{ttl}: "+'; '.join(cl[:6])+(' …' if len(cl)>6 else ''))
        ft=('Eixo ajustado aos valores da Dexco.' if gid=='foco' else '')
        if notes: ft+=' Fora do eixo: '+' | '.join(notes)
        if exnotes: ft+=' Excluídos (ver Ajustes na base): '+' | '.join(exnotes)
        if any(s['metrics'].get(k,{}).get('calc') for k,_ in [(i[1],0) for i in items if i[0]=='M']): ft+=' (1) EBITDA = ROL × margem EBITDA.'
        footnote(sl,ft[:900])
# ---------- ajustes na base ----------
def short(w):
    for k,v in [('lucro líquido igual','Lucro = receita (célula puxa a linha errada)'),('margem de 100% decorrente','Margem de 100% (mesmo erro do lucro)'),('margem de exatamente','Margem de exatamente 100%'),('prazo médio negativo','Prazo negativo (impossível)'),
                ('dívida líquida/EBITDA','|DL/EBITDA| > 10x (EBITDA ≈ 0)'),('acima de 500%','|Índice| > 500% (EBITDA ≈ 0)'),('prazo acima de 365','Prazo > 365 dias'),('giro acima de 6x','Giro > 6x (CI muito baixo)'),
                ('margem abaixo','Margem < −50%'),('ROIC abaixo','ROIC < −50%'),('capital de giro abaixo','Capital de giro < −50% da ROL'),('capex acima','Capex > 50% da ROL (implantação)')]:
        if k in w: return v
    return w
for seg in ['loucas','metais','madeira','revest']:
    s=D[seg];log=s.get('log') or []
    if not log: continue
    sl=new_slide();SEGSLIDES[seg]['ajustes']=sl
    ne=sum(1 for l in log if l['kind']=='erro')
    title(sl,f"Peers {s['name']} – Ajustes na base",f"{ne} erro(s) de origem e {len(log)-ne} outlier(s) retirados dos gráficos, da matriz e das medianas (aparecem como n.s.)")
    cols=[('Tipo',.7),('Indicador',2.0),('Empresa',1.7),('Ano',.55),('Valor',1.0),('Motivo',3.35),('Célula na planilha',2.75)]
    rows=len(log)+1;rh=min(.26,(6.95-1.15)/rows)
    gt=sl.shapes.add_table(rows,len(cols),Inches(1.05),Inches(1.15),Inches(sum(c[1] for c in cols)),Inches(rh*rows)).table
    for j,(c,wd) in enumerate(cols): gt.columns[j].width=Inches(wd)
    def cellset(cell,text,bold=False,color='262626',fill=None,align=PP_ALIGN.LEFT):
        cell.text='';p=cell.text_frame.paragraphs[0];p.alignment=align;r=p.add_run();r.text=text;r.font.size=Pt(8 if rows>18 else 9);r.font.bold=bold;r.font.name='Roboto';r.font.color.rgb=RGB(color)
        cell.margin_top=cell.margin_bottom=Inches(.01);cell.margin_left=cell.margin_right=Inches(.05);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        cell.fill.solid();cell.fill.fore_color.rgb=RGB(fill or 'FFFFFF')
    for j,(c,_) in enumerate(cols): cellset(gt.cell(0,j),c,True,'FFFFFF','404040')
    for i,l in enumerate(log,1):
        vals=['Erro' if l['kind']=='erro' else 'Outlier',l['metric'],l['ent'],l['year'],fmt(l['v'],l['fmt'],False),short(l['why']),l['cell']]
        for j,v in enumerate(vals): cellset(gt.cell(i,j),v,j==0 and l['kind']=='erro',fill='F2F2F2' if i%2==0 else 'FFFFFF',align=PP_ALIGN.RIGHT if j in (3,4) else PP_ALIGN.LEFT)
    for i in range(rows): gt.rows[i].height=Inches(rh)
    footnote(sl,'Erro: o valor não pode estar certo (ex.: lucro igual à receita, prazo negativo). Outlier: o valor pode estar certo na planilha, mas o índice perde sentido (ex.: dívida/EBITDA com EBITDA perto de zero). Regras: |DL/EBITDA| > 10x; |índices sobre EBITDA| > 500%; prazos > 365 dias; giro > 6x; margens e ROIC < −50%; capex > 50% da ROL.')
# ---------- navegação ----------
def add_nav(sl,seg,cur):
    x=13.1;btns=[]
    for key,lab in reversed(NAV):
        if key not in ('menu',) and key not in SEGSLIDES[seg]: continue
        wd=.34+len(lab)*.06;x-=wd+.05
        b=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(.64),Inches(wd),Inches(.26))
        on=key==cur;b.fill.solid();b.fill.fore_color.rgb=RGB(SEGCOL[seg] if on else 'F2F2F2');b.line.fill.background();b.shadow.inherit=False
        tf=b.text_frame;tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=lab;r.font.size=Pt(8.5);r.font.bold=on;r.font.name='Roboto';r.font.color.rgb=RGB('FFFFFF' if on else '404040')
        target=menu if key=='menu' else SEGSLIDES[seg][key]
        if not on: b.click_action.target_slide=target
    # etiqueta do negócio
for seg,d in SEGSLIDES.items():
    for key,sl in d.items(): add_nav(sl,seg,key)
# menu: botões levam ao resumo de cada negócio
for nm,seg in [('Rectangle 5','loucas'),('Rectangle 7','metais'),('Rectangle 9','madeira'),('Rectangle 11','revest')]:
    find(menu,nm).click_action.target_slide=summary[seg]
# ---------- ordem dos slides ----------
lst=prs.slides._sldIdLst;ids=list(lst)
byslide={sl.slide_id:el for sl,el in zip(prs.slides,ids)}
order=[cover,menu]
for seg in ['loucas','metais','madeira','revest']:
    order+= [SEGSLIDES[seg][g] for g in ['resumo','escala','foco','rent','custos','giro','caixa','setor','ajustes'] if g in SEGSLIDES[seg]]
order.append(closing)
els=[byslide[sl.slide_id] for sl in order]
for el in ids: lst.remove(el)
for el in els: lst.append(el)
out=sys.argv[3];prs.save(out);print('saved',len(order),'slides')

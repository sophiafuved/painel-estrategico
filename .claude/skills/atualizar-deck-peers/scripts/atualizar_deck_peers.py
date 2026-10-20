"""Atualiza um PowerPoint de Análise de Peers (Louças, Metais, Madeira ou Revestimento) com os dados atuais
da planilha Análise de Competidores, sem quebrar o vínculo dos gráficos e mantendo a formatação.

Por que reconhecer por conteúdo: os gráficos dos decks apontam para células da planilha (ex.:
'Gráficos - Louças'!$D$67:$J$67), mas a planilha ganha linhas e blocos novos com o tempo e os endereços
antigos passam a apontar para outro dado. O script reconhece cada série pelo **nome da empresa e pelos
valores** e encontra a linha certa na planilha atual:
  1. Série com o mesmo nome e os mesmos valores (≥ 60% dos pontos) numa linha de um bloco → essa linha.
  2. Demais séries do mesmo gráfico (valores revisados, série sem nome) → mesmo deslocamento de linhas das
     séries já reconhecidas, se o nome da linha bater.
  3. O que não for reconhecido fica como estava e vai para a conferência.
Depois, regrava os valores e as faixas (vínculo continua com a planilha), acrescenta anos novos quando o
bloco da planilha já tem dado para eles e gera uma conferência em Markdown.

Uso:
  python atualizar_deck_peers.py --projeto <painel-estrategico> --planilha "<Análise de Competidores>.xlsx" \\
      --deck "<Análise de Peers - Louças>.pptx" --saida "Analise_de_Peers_Loucas_atualizado.pptx"
"""
import argparse, collections, json, os, re, subprocess, sys, tempfile, unicodedata, warnings
warnings.filterwarnings('ignore')
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from pptx import Presentation

C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
q = lambda t: f'{{{C}}}{t}'
REF = re.compile(r"^(?P<pre>'?(?:\[[^\]]*\])?(?P<sheet>[^'!\[]+)'?)!\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$")
SHEETS = ['Gráficos - Louças', 'Gráficos - Metais', 'Gráficos Madeira 1', 'Gráficos - Revestimento']
NEGOCIO = dict(zip(SHEETS, ['Louças', 'Metais', 'Madeira', 'Revestimento']))


def norm(t):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower())


def isyear(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 2000 <= v <= 2035 and float(v).is_integer():
        return True
    return isinstance(v, str) and bool(re.fullmatch(r'\s*(20\d\d|\d?T\d\d.*|.*LTM.*|20\d\dE)\s*', v))


def ystr(v):
    if isinstance(v, (int, float)):
        return str(int(v))
    s = re.sub(r'\s*LTM\s*', '', str(v)).strip()  # "1T26 LTM" e "1T26" são o mesmo período
    return str(int(float(s))) if re.fullmatch(r'\d{4}(\.0+)?', s) else s


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def igual(a, b):
    return a is not None and b is not None and abs(a - b) <= 1e-6 * max(1, abs(a))


# ---------- caches dos gráficos ----------
def read_cache(cache):
    if cache is None:
        return []
    pc = cache.find(q('ptCount'))
    n = int(pc.get('val')) if pc is not None else 0
    out = [None] * n
    for pt in cache.findall(q('pt')):
        i = int(pt.get('idx'))
        if i < n:
            t = pt.find(q('v')).text
            try:
                out[i] = float(t)
            except (TypeError, ValueError):
                out[i] = t
    return out


def write_cache(cache, vals, text=False):
    for pt in cache.findall(q('pt')):
        cache.remove(pt)
    pc = cache.find(q('ptCount'))
    if pc is None:
        pc = etree.SubElement(cache, q('ptCount'))
    pc.set('val', str(len(vals)))
    anchor = cache.find(q('extLst'))
    for i, v in enumerate(vals):
        if v is None:
            continue
        pt = etree.Element(q('pt'), idx=str(i))
        etree.SubElement(pt, q('v')).text = (str(v) if text else repr(float(v)))
        (anchor.addprevious(pt) if anchor is not None else cache.append(pt))


def cache_of(el):
    if el is None:
        return None
    return el.find('.//' + q('numCache')) if el.find('.//' + q('numCache')) is not None else el.find('.//' + q('strCache'))


# ---------- planilha ----------
def indexar(wb):
    """Linhas de empresa de cada bloco das abas "Gráficos": {aba: [dict(row, name, title, hrow, ncol, yrs={ano: col})]}."""
    idx = {}
    for sh in SHEETS:
        if sh not in wb.sheetnames:
            continue
        ws = wb[sh]; out = []; hdr = None
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            row = list(row) + [None] * 45
            achou = None
            for c in range(0, 30):
                if isinstance(row[c], str) and row[c].strip() and not isyear(row[c]) and isyear(row[c + 1]) and isyear(row[c + 2]):
                    achou = c; break
            if achou is not None:
                yrs = {}; k = achou + 1
                while isyear(row[k]):
                    yrs[ystr(row[k])] = k + 1; k += 1
                hdr = dict(title=row[achou].strip(), hrow=r, ncol=achou + 1, yrs=yrs)
                continue
            # linhas abaixo do cabeçalho, inclusive sub-blocos separados por linha em branco (ex.: participação
            # logo abaixo dos valores nos blocos "Mkt" de Metais); o casamento por valores evita confusão
            if hdr and row[hdr['ncol'] - 1] not in (None, ''):
                pais = row[hdr['ncol'] - 2] if hdr['ncol'] >= 2 else None
                out.append(dict(hdr, row=r, name=str(row[hdr['ncol'] - 1]).strip(),
                                pais=(pais.strip() if isinstance(pais, str) and 1 < len(pais.strip()) <= 4 else None)))
        idx[sh] = out
    return idx


def ajustes_do_painel(projeto, planilha):
    """Valores que o painel trata como erro/outlier/n.s. (clean.py), por célula de origem."""
    pp = os.path.abspath(os.path.join(projeto, 'pipeline', 'peers'))
    if not os.path.exists(os.path.join(pp, 'clean.py')):
        print(f'Aviso: não encontrei {pp}/clean.py; a conferência sai sem os alertas de n.s./outlier.', file=sys.stderr)
        return {}
    w = tempfile.mkdtemp(); env = {**os.environ, 'PEERS_XLSX': os.path.abspath(planilha)}
    for script, args in (('parse_blocks.py', SHEETS), ('build.py', []), ('clean.py', [])):
        r = subprocess.run([sys.executable, '-I', os.path.join(pp, script), *args], cwd=w, env=env, capture_output=True, text=True)
        if r.returncode != 0:
            print(f'Aviso: não consegui rodar {script} do painel; a conferência sai sem os alertas de n.s./outlier.', file=sys.stderr)
            return {}
    out = {}
    for a in json.load(open(os.path.join(w, 'ajustes.json'), encoding='utf-8')):
        out[a['cell'].replace('$', '')] = a
    return out


# ---------- séries ----------
def info_serie(ser, wb):
    val = ser.find(q('val')) if ser.find(q('val')) is not None else ser.find(q('yVal'))
    cat = ser.find(q('cat')) if ser.find(q('cat')) is not None else ser.find(q('xVal'))
    tx = ser.find(q('tx'))
    vf = val.find('.//' + q('f')) if val is not None else None
    m = REF.match(vf.text or '') if vf is not None else None
    nome = (read_cache(cache_of(tx)) or [''])[0] if tx is not None else ''
    cats = read_cache(cache_of(cat)) if cat is not None else []
    vals = read_cache(cache_of(val)) if val is not None else []
    return dict(ser=ser, val=val, cat=cat, tx=tx, vf=vf, m=m, nome=str(nome or '').strip(),
                anos=[ystr(c) if c is not None else None for c in cats], vals=[v if isinstance(v, float) else None for v in vals],
                aba=(m.group('sheet') if m else None), na_planilha=bool(m and m.group('sheet') in wb.sheetnames and m.group('sheet') in SHEETS))


def casar(s, linhas):
    """Melhor linha da planilha para a série: mesmo nome (se houver) e mais valores iguais nos mesmos anos."""
    melhor = None
    pontos = [(i, v) for i, v in enumerate(s['vals']) if v]  # zeros não servem de evidência (há zeros em todo bloco)
    for L in linhas:
        if s['nome'] and norm(L['name']) != norm(s['nome']):
            continue
        ok = 0
        for i, v in pontos:
            a = s['anos'][i] if i < len(s['anos']) else None
            if a in L['yrs'] and igual(num(L['_ws'].cell(L['row'], L['yrs'][a]).value), v):
                ok += 1
        if ok and (melhor is None or ok > melhor[0]):
            melhor = (ok, L)
    if melhor and melhor[0] >= max(1, 0.6 * len(pontos)):
        return melhor[1]
    if s['nome'] and len(pontos) >= 3 and len({round(v, 9) for _, v in pontos}) >= 3:
        # nome diferente na planilha (ex.: "Deca" × "Deca Louças"): aceita só se todos os valores baterem
        cand = [L for L in linhas if all((s['anos'][i] if i < len(s['anos']) else None) in L['yrs']
                                         and igual(num(L['_ws'].cell(L['row'], L['yrs'][s['anos'][i]]).value), v) for i, v in pontos)]
        if len(cand) == 1:
            return cand[0]
    return None


def regravar(s, L, ws, anos, log_item, ajustes):
    """Aponta a série para a linha L nos anos `anos` e regrava valores, categorias e nome."""
    cols = [L['yrs'][a] for a in anos]
    novos = [num(ws.cell(L['row'], c).value) for c in cols]
    pre = s['m'].group('pre')
    contiguo = cols == list(range(cols[0], cols[0] + len(cols)))
    antes = s['vf'].text
    if contiguo:
        s['vf'].text = f"{pre}!${gl(cols[0])}${L['row']}:${gl(cols[-1])}${L['row']}"
    else:
        s['vf'].text = f"{pre}!" + ','.join(f"${gl(c)}${L['row']}" for c in cols)
    write_cache(cache_of(s['val']), novos)
    log_item['revisoes'] = sum(1 for o, n in zip(s['vals'], novos) if o is not None and n is not None and not igual(o, n))
    log_item['sumiram'] = sum(1 for o, n in zip(s['vals'], novos) if o is not None and n is None)
    log_item['faixa_antes'], log_item['faixa_depois'] = antes, s['vf'].text
    # categorias = linha de anos do bloco
    if s['cat'] is not None:
        cf = s['cat'].find('.//' + q('f'))
        if cf is not None and contiguo:
            cm = REF.match(cf.text or '')
            if cm:
                cf.text = f"{cm.group('pre')}!${gl(cols[0])}${L['hrow']}:${gl(cols[-1])}${L['hrow']}"
        cc = cache_of(s['cat'])
        if cc is not None:
            if cc.tag == q('numCache'):
                write_cache(cc, [float(a) if re.fullmatch(r'\d{4}', a) else None for a in anos])
            else:
                write_cache(cc, anos, text=True)
    # nome da série vinculado à planilha
    if s['tx'] is not None:
        tf = s['tx'].find('.//' + q('f')); tm = REF.match(tf.text or '') if tf is not None else None
        if tm and tm.group('sheet') == L['_aba']:
            tf.text = f"{tm.group('pre')}!${gl(L['ncol'])}${L['row']}"
            sc = cache_of(s['tx'])
            if sc is not None:
                write_cache(sc, [L['name']], text=True)
    # rótulo do último ponto acompanha o dado novo
    li_old = max((i for i, v in enumerate(s['vals']) if v is not None), default=None)
    li_new = max((i for i, v in enumerate(novos) if v is not None), default=None)
    dl = s['ser'].find(q('dLbls'))
    if dl is not None and li_old is not None and li_new is not None and li_new != li_old:
        for d in dl.findall(q('dLbl')):
            ix = d.find(q('idx'))
            if ix is not None and int(ix.get('val')) == li_old:
                ix.set('val', str(li_new)); log_item['rotulo_movido'] = True
    for a, c in zip(anos, cols):
        aj = ajustes.get(f"'{L['_aba']}'!{gl(c)}{L['row']}")
        if aj:
            log_item.setdefault('alertas', []).append(f"{a}: {aj['kind']} no painel ({aj['why']})")


def celula_unica(s, wb, L):
    """Série de uma célula só (ex.: bolhas/dispersão): relê o valor no mesmo endereço."""
    m = s['m']; v = num(wb[m.group('sheet')][f"{m.group('c1')}{m.group('r1')}"].value)
    s['_unica'] = (v, s['vals'][0] if s['vals'] else None)
    write_cache(cache_of(s['val']), [v])


def endereco_confere(s, wb):
    """O endereço antigo da série ainda traz os valores que estão no gráfico?"""
    m = s['m']; ws = wb[m.group('sheet')]
    c1, r1 = ci(m.group('c1')), int(m.group('r1'))
    c2, r2 = ci(m.group('c2') or m.group('c1')), int(m.group('r2') or m.group('r1'))
    cel = [num(ws.cell(r, c).value) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    pts = [(i, v) for i, v in enumerate(s['vals']) if v is not None]
    return bool(pts) and all(i < len(cel) and igual(cel[i], v) for i, v in pts)


def fixar(s, tirar=()):
    """Troca as referências da série por valores fixos (numLit/strLit), opcionalmente tirando pontos."""
    def lit(ref, tag):
        cache = ref.find(q('numCache')) if ref.find(q('numCache')) is not None else ref.find(q('strCache'))
        if cache is None:
            return
        if tirar:
            vals = read_cache(cache)
            vals = [v for i, v in enumerate(vals) if i not in tirar]
            write_cache(cache, vals, text=cache.tag == q('strCache'))
        cache.tag = q('numLit') if cache.tag == q('numCache') else q('strLit')
        ref.addprevious(cache); ref.getparent().remove(ref)
    for el in (s['val'], s['cat']):
        if el is None:
            continue
        for ref in list(el):
            if ref.tag in (q('numRef'), q('strRef')):
                lit(ref, ref.tag)
    if s['tx'] is not None:
        ref = s['tx'].find(q('strRef'))
        if ref is not None:
            nome = (read_cache(ref.find(q('strCache'))) or [''])[0]
            v = etree.Element(q('v')); v.text = str(nome or '')
            ref.addprevious(v); s['tx'].remove(ref)


def atualizar(prs, wb, idx, ajustes, ano_max):
    log, sem_vinculo = [], []
    for sh, linhas in idx.items():
        for L in linhas:
            L['_ws'] = wb[sh]; L['_aba'] = sh
    for sn, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not getattr(shp, 'has_chart', False):
                continue
            series = [info_serie(s, wb) for s in shp.chart._chartSpace.iter(q('ser'))]
            casadas = {}
            for k, s in enumerate(series):
                if s['na_planilha'] and s['m'].group('r2') and s['m'].group('r1') == s['m'].group('r2'):
                    L = casar(s, idx[s['aba']])
                    if L:
                        casadas[k] = (L, 'nome e valores')
            # mesmo endereço e mesmo nome na linha (a planilha não mudou de lugar, só os valores foram revisados)
            for k, s in enumerate(series):
                if k in casadas or not s['na_planilha'] or not s['nome'] or s['m'].group('r1') != (s['m'].group('r2') or s['m'].group('r1')):
                    continue
                L = next((x for x in idx[s['aba']] if x['row'] == int(s['m'].group('r1'))), None)
                if L and norm(L['name']) == norm(s['nome']):
                    if s['m'].group('c2') and s['m'].group('c2') != s['m'].group('c1'):
                        casadas[k] = (L, 'mesmo endereço e nome (valores revisados)')
                    else:
                        celula_unica(s, wb, L)
            # nenhuma série reconhecida (ex.: todos os valores revisados) e a planilha ganhou linhas: procura um
            # deslocamento pequeno em que TODAS as séries com nome encontrem a linha com o mesmo nome
            linhas_ser = [k for k, s in enumerate(series) if s['na_planilha'] and s['nome'] and s['m'].group('r2')
                          and s['m'].group('r1') == s['m'].group('r2') and s['m'].group('c2') != s['m'].group('c1')]
            if not casadas and len(linhas_ser) >= 2:
                for d in sorted(range(-15, 16), key=abs):
                    achou = {}
                    for k in linhas_ser:
                        s = series[k]
                        L = next((x for x in idx[s['aba']] if x['row'] == int(s['m'].group('r1')) + d), None)
                        if not (L and norm(L['name']) == norm(s['nome'])):
                            break
                        achou[k] = (L, f'nome da empresa {"no mesmo endereço" if d == 0 else f"{d:+d} linhas"} (valores revisados)')
                    else:
                        casadas.update(achou); break
            # deslocamento de linhas das séries reconhecidas, para as demais do mesmo gráfico
            # (só quando metade ou mais das séries do gráfico foi reconhecida por conteúdo e todas concordam)
            desl = collections.Counter(L['row'] - int(series[k]['m'].group('r1')) for k, (L, _) in casadas.items())
            na_pl = sum(1 for s in series if s['na_planilha'])
            confiavel = len(desl) == 1 and len(casadas) >= max(1, na_pl / 2)
            for k, s in enumerate(series):
                # com nome: basta o deslocamento ser único e o nome bater; sem nome: exige a maioria reconhecida
                if k in casadas or not s['na_planilha'] or len(desl) != 1 or (not s['nome'] and not confiavel):
                    continue
                d = desl.most_common(1)[0][0]
                alvo = int(s['m'].group('r1')) + d
                L = next((x for x in idx[s['aba']] if x['row'] == alvo), None)
                if L and (not s['nome'] or norm(L['name']) == norm(s['nome'])) and all(a in L['yrs'] for a in s['anos'] if a):
                    casadas[k] = (L, 'posição (valores revisados ou série sem nome)')
            # anos do gráfico: os da série mais longa que existem em todos os blocos casados, mais os anos seguintes
            # do bloco que já têm dado para alguma série; períodos que a planilha não tem mais (ex.: 1T26) saem
            anos, saiu, base = None, [], []
            if casadas:
                base = max((series[k]['anos'] for k in casadas), key=len)
                Ls = [L for L, _ in casadas.values()]
                if base and all(base):
                    anos = [a for a in base if all(a in L['yrs'] for L in Ls)]
                    saiu = [a for a in base if a not in anos]
                    todos = [a for a in Ls[0]['yrs'] if all(a in L['yrs'] for L in Ls)]
                    if anos and anos[-1] in todos and not saiu:
                        for a in todos[todos.index(anos[-1]) + 1:]:
                            if re.fullmatch(r'20\d\d', a) and int(a) <= ano_max and any(num(L['_ws'].cell(L['row'], L['yrs'][a]).value) is not None for L in Ls):
                                anos.append(a)
                            else:
                                break
                    if saiu and any(not re.search(r'T\d\d', a) for a in saiu):
                        anos = None  # falta um ano inteiro no bloco: melhor não mexer
            for k, s in enumerate(series):
                item = dict(slide=sn, grafico=shp.name, serie=s['nome'] or f'série {k + 1}', faixa_antes=s['vf'].text if s['vf'] is not None else '',
                            faixa_depois='', bloco='', como='', revisoes=0, sumiram=0, anos_novos=[], rotulo_movido=False, obs='')
                L, como = casadas.get(k, (None, ''))
                curta = bool(L and anos) and len(s['anos']) < len(base)  # série mais curta no deck continua do mesmo tamanho
                anos_s = ([a for a in s['anos'] if a in anos] or anos[:len(s['vals'])]) if curta else (anos or [])
                anos_s = [a for a in anos_s if L and a in L['yrs']]
                if L and anos_s:
                    item['bloco'] = f"{L['title']} (linha {L['row']})"; item['como'] = como
                    item['anos_novos'] = [a for a in anos_s if a not in (s['anos'] if any(s['anos']) else base)]
                    if saiu:
                        item['obs'] = f"a planilha não tem mais {', '.join(saiu)} neste bloco: período retirado do gráfico"
                    regravar(s, L, wb[L['_aba']], anos_s, item, ajustes)
                elif s.get('_unica'):
                    v, o = s['_unica']
                    item['faixa_depois'] = item['faixa_antes']; item['como'] = 'mesmo endereço e nome (célula única)'
                    item['revisoes'] = int(o is not None and v is not None and not igual(o, v)); item['sumiram'] = int(o is not None and v is None)
                elif s['na_planilha'] and endereco_confere(s, wb):
                    item['obs'] = 'mantida: o endereço antigo ainda traz os mesmos valores (vínculo preservado)'
                elif s['na_planilha']:
                    # o endereço antigo pode apontar hoje para outra linha (a planilha ganhou linhas): fixa os valores
                    # no gráfico, sem vínculo, para o "Atualizar links" não trocar por dado de outra empresa
                    fixar(s, [i for i, a in enumerate(s['anos']) if a in saiu])
                    item['obs'] = ('não reconhecida na planilha atual: valores antigos fixos no gráfico, sem vínculo '
                                   '(a linha pode ter sido apagada ou renomeada; confira e apague a série se não fizer mais sentido)')
                elif s['aba']:
                    item['obs'] = f'dados de "{s["aba"]}" (outra planilha ou dados colados no gráfico): mantidos'
                    sem_vinculo.append(item)
                else:
                    item['obs'] = 'dados colados no gráfico (sem vínculo): mantidos'
                    sem_vinculo.append(item)
                log.append(item)
    return log, sem_vinculo


# ---------- textos para revisar ----------
def textos_para_revisar(prs, idx, negocio, pend):
    outros = set()
    for sh, linhas in idx.items():
        if NEGOCIO.get(sh) != negocio:
            outros |= {L['name'] for L in linhas if L.get('pais')}  # só empresas (linhas com país ao lado)
    deste = {norm(L['name']) for sh, linhas in idx.items() if NEGOCIO.get(sh) == negocio for L in linhas}
    alheios = {n for n in outros if len(n) > 4 and norm(n) not in deste and not norm(n).startswith('dexco')}
    for sn, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            t = shp.text_frame.text.strip()
            if not t:
                continue
            if re.search(r'\[WIP\]|lorem ipsum', t, re.I) or t.lower() == 'x':
                pend.append(f'Slide {sn}: texto provisório "{t[:60]}".')
                continue
            cit = sorted({n for n in alheios if re.search(r'\b' + re.escape(n) + r'\b', t)})
            if cit:
                pend.append(f'Slide {sn}: o texto cita {", ".join(cit)}, que não está(ão) na aba de {negocio} da planilha (texto copiado de outro deck?): "{t[:80]}…"')


def conferencia(path, log, sem_vinculo, pend, args, negocio, vinc):
    ok = [x for x in log if x['faixa_depois']]
    rev = [x for x in ok if x['revisoes'] or x['sumiram']]
    novos = sorted({a for x in ok for a in x['anos_novos']})
    nao = [x for x in log if x['obs'].startswith('não reconhecida')]
    mant = [x for x in log if x['obs'].startswith('mantida: o endereço')]
    L = [f'# Conferência do deck de peers: {negocio}', '',
         f'- Deck: `{os.path.basename(args.deck)}` · Planilha: `{os.path.basename(args.planilha)}`',
         f'- Vínculo dos gráficos no deck: {", ".join(f"`{v}`" for v in vinc) or "nenhum"}', '',
         '## Resumo',
         f'- Séries ligadas às abas "Gráficos": {len(ok) + len(nao) + len(mant)} · atualizadas: {len(ok)} · endereço antigo ainda certo (mantidas): {len(mant)} · não reconhecidas (valores fixos): {len(nao)}',
         f'- Séries com valores revisados na planilha: {len(rev)}',
         f'- Anos acrescentados: {", ".join(novos) if novos else "nenhum"}',
         f'- Séries sem vínculo com a planilha (dados colados ou outro arquivo): {len(sem_vinculo)}', '',
         '## Revisar à mão'] + [f'- {p}' for p in pend]
    for x in log:
        if x.get('alertas'):
            L.append(f"- Slide {x['slide']} · {x['grafico']} · {x['serie']}: " + '; '.join(x['alertas']) + ' (o deck mostra o valor da planilha)')
    saidas = collections.OrderedDict()
    for x in ok:
        if x['obs']:
            saidas.setdefault((x['slide'], x['grafico'], x['obs']), 0)
    for (sl, g, o) in saidas:
        L.append(f'- Slide {sl} · {g}: {o}.')
    for x in nao:
        L.append(f"- Slide {x['slide']} · {x['grafico']} · {x['serie']}: {x['obs']} (antes: {x['faixa_antes']})")
    if sem_vinculo:
        por = collections.Counter(x['slide'] for x in sem_vinculo)
        L.append('- Gráficos com dados colados ou de outra planilha (não atualizados): ' + ', '.join(f'slide {s} ({n} séries)' for s, n in sorted(por.items())))
    L += ['', '## Séries com valores revisados', '| Slide | Gráfico | Série | Bloco | Pontos alterados | Pontos que ficaram vazios |', '|---|---|---|---|---|---|']
    L += [f"| {x['slide']} | {x['grafico']} | {x['serie']} | {x['bloco']} | {x['revisoes']} | {x['sumiram']} |" for x in rev]
    L += ['', '## Todas as séries', '| Slide | Gráfico | Série | Bloco | Reconhecida por | Faixa |', '|---|---|---|---|---|---|']
    for x in log:
        fx = x['faixa_depois'] if x['faixa_antes'] == x['faixa_depois'] else f"{x['faixa_antes']} → {x['faixa_depois'] or ('valores fixos' if x['obs'].startswith('não reconhecida') else 'mantida')}"
        L.append(f"| {x['slide']} | {x['grafico']} | {x['serie']} | {x['bloco'] or '–'} | {x['como'] or x['obs']} | {fx} |")
    open(path, 'w', encoding='utf-8').write('\n'.join(L) + '\n')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--projeto', default='.'); ap.add_argument('--planilha', required=True)
    ap.add_argument('--deck', required=True); ap.add_argument('--saida', required=True); ap.add_argument('--conferencia')
    ap.add_argument('--ano-max', type=int, default=None, help='último ano a acrescentar (padrão: ano passado)')
    args = ap.parse_args()
    import datetime
    ano_max = args.ano_max or datetime.date.today().year - 1
    wb = load_workbook(args.planilha, data_only=True)
    idx = indexar(wb)
    ajustes = ajustes_do_painel(args.projeto, args.planilha)
    prs = Presentation(args.deck)
    # negócio do deck = aba mais usada pelos gráficos
    abas = collections.Counter()
    vinc = set()
    for slide in prs.slides:
        for shp in slide.shapes:
            if getattr(shp, 'has_chart', False):
                for f in shp.chart._chartSpace.iter(q('f')):
                    m = REF.match(f.text or '')
                    if m and m.group('sheet') in SHEETS:
                        abas[m.group('sheet')] += 1
                for rel in shp.chart.part.rels.values():
                    if rel.is_external:
                        vinc.add(rel.target_ref.split('/')[-1].replace('%20', ' '))
    negocio = NEGOCIO.get(abas.most_common(1)[0][0], '?') if abas else '?'
    pend = []
    log, sem_vinculo = atualizar(prs, wb, idx, ajustes, ano_max)
    textos_para_revisar(prs, idx, negocio, pend)
    prs.save(args.saida)
    conf = args.conferencia or os.path.splitext(args.saida)[0] + '_conferencia.md'
    conferencia(conf, log, sem_vinculo, pend, args, negocio, sorted(vinc))
    ok = [x for x in log if x['faixa_depois']]
    print(json.dumps({'negocio': negocio, 'saida': os.path.abspath(args.saida), 'conferencia': os.path.abspath(conf),
                      'series_na_planilha': len(ok) + sum(1 for x in log if x['obs'].startswith('não reconhecida')),
                      'atualizadas': len(ok), 'com_revisao': sum(1 for x in ok if x['revisoes'] or x['sumiram']),
                      'anos_novos': sorted({a for x in ok for a in x['anos_novos']}), 'sem_vinculo': len(sem_vinculo),
                      'alertas_painel': sum(len(x.get('alertas', [])) for x in log), 'revisar_a_mao': len(pend), 'vinculo': sorted(vinc)},
                     ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()

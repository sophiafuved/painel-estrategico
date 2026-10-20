"""Leva um deck do time para o template PowerPoint da Dexco (padrão: assets/Template_PPT_2026.potx).

Monta um deck novo a partir do template e copia o conteúdo de cada slide do deck original, mantendo
gráficos (com o vínculo à planilha), tabelas, imagens, caixas de texto e hyperlinks entre slides:
  - capa (1º slide)            → capa do template: TÍTULO / Sub-título / | data |
  - slide vazio de encerramento → encerramento do template (logos Dexco)
  - divisor de seção ("Capa de início" no meio do deck) → agenda do template (seção atual em destaque)
  - demais slides               → slide de conteúdo do template: etiqueta cinza no topo + título 20 pt;
    o título antigo vira o título novo e o restante é copiado na mesma posição
Gera também uma conferência em Markdown (títulos, slides com conteúdo invadindo a faixa do título).

Uso:
  python aplicar_template.py --deck "<deck>.pptx" --saida "<deck>_template2026.pptx" \\
      [--template Template_PPT_2026.potx] [--etiqueta "Análise de Peers: Louças"] [--data "Offsite 2026"]
"""
import argparse, copy, io, os, re, sys, zipfile
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Inches

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
qa = lambda t: f'{{{A}}}{t}'
qp = lambda t: f'{{{P}}}{t}'
AQUI = os.path.dirname(os.path.abspath(__file__))
PADRAO = os.path.join(AQUI, '..', 'assets', 'Template_PPT_2026.potx')
TOPO_CONTEUDO = Inches(1.18)  # o título do template vai até ~1,18 pol.


# ---------- abrir o template (.potx) ----------
def abrir_template(path):
    if path.lower().endswith('.pptx'):
        return Presentation(path)
    zin = zipfile.ZipFile(path); buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == '[Content_Types].xml':
                data = data.replace(b'presentationml.template.main+xml', b'presentationml.presentation.main+xml')
            zout.writestr(it, data)
    buf.seek(0)
    return Presentation(buf)


# ---------- cópia de formas com as relações ----------
class Copiador:
    def __init__(self, dst):
        self.dst = dst
        self.usados = {str(p.partname) for p in dst.part.package.iter_parts()}
        self.renomeados = set()
        self.links_slide = []  # (parte destino, elemento, atributo, parte do slide antigo, tipo)

    def _nome_livre(self, partname):
        base, ext = os.path.splitext(str(partname))
        base = re.sub(r'\d+$', '', base)
        i = 1
        while f'{base}{i}{ext}' in self.usados:
            i += 1
        nome = f'{base}{i}{ext}'
        self.usados.add(nome)
        return PackURI(nome)

    def _renomear(self, part):
        """Dá nomes únicos à parte e às partes ligadas a ela (gráfico → planilha embutida, estilo, cores...)."""
        if id(part) in self.renomeados:
            return
        self.renomeados.add(id(part))
        part.partname = self._nome_livre(part.partname)
        for rel in part.rels.values():
            if not rel.is_external:
                self._renomear(rel.target_part)

    def copiar(self, el_src, src_part, dst_slide):
        el = copy.deepcopy(el_src)
        dst_part = dst_slide.part
        for node in el.iter():
            for att in list(node.attrib):
                if not att.startswith(f'{{{R}}}'):
                    continue
                rid = node.get(att)
                if rid not in src_part.rels:
                    continue
                rel = src_part.rels[rid]
                if rel.is_external:
                    node.set(att, dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True))
                elif rel.reltype == RT.SLIDE:
                    self.links_slide.append((dst_part, node, att, rel.target_part, rel.reltype))
                else:
                    self._renomear(rel.target_part)
                    node.set(att, dst_part.relate_to(rel.target_part, rel.reltype))
        # ids únicos no slide (o PowerPoint pede reparo se dois objetos tiverem o mesmo id)
        tree = dst_slide.shapes._spTree
        usados = {int(c.get('id')) for c in tree.iter(qp('cNvPr')) if (c.get('id') or '').isdigit()}
        prox = max(usados | {1}) + 1
        for c in el.iter(qp('cNvPr')):
            c.set('id', str(prox)); prox += 1
        tree.append(el)
        return el

    def resolver_links(self, mapa):
        for dst_part, node, att, alvo, reltype in self.links_slide:
            novo = mapa.get(id(alvo))
            if novo is not None:
                node.set(att, dst_part.relate_to(novo.part, reltype))
            else:  # link para slide que não existe mais: tira o link
                node.getparent().remove(node) if node.tag.endswith('hlinkClick') else node.attrib.pop(att, None)


def sem_placeholder(el, shape):
    """Placeholder do layout antigo vira forma comum, com a posição que tinha (o layout novo é outro)."""
    nv = el.find('.//' + qp('nvPr'))
    ph = nv.find(qp('ph')) if nv is not None else None
    if ph is None:
        return
    nv.remove(ph)
    spPr = el.find(qp('spPr'))
    if spPr is not None and spPr.find(qa('xfrm')) is None:
        x = etree.SubElement(spPr, qa('xfrm')); spPr.insert(0, x)
        etree.SubElement(x, qa('off'), x=str(shape.left or 0), y=str(shape.top or 0))
        etree.SubElement(x, qa('ext'), cx=str(shape.width or 0), cy=str(shape.height or 0))


# ---------- classificação dos slides antigos ----------
def eh_thinkcell(sh):
    return 'think-cell' in sh.name.lower() or 'thinkcell' in sh.name.lower()


def titulo_antigo(slide):
    cand = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.top is None:
            continue
        ph = sh.placeholder_format.type if sh.is_placeholder else None
        if (ph is not None and 'TITLE' in str(ph)) or (sh.top < Inches(0.75) and sh.width > Inches(4) and sh.height < Inches(1.1)):
            cand.append(sh)
    return min(cand, key=lambda s: (s.top, s.left)) if cand else None


def tipo_slide(i, slide):
    formas = [s for s in slide.shapes if not eh_thinkcell(s)]
    if i == 0:
        return 'capa'
    if not formas:
        return 'encerramento'
    if re.search(r'capa|in[ií]cio', slide.slide_layout.name, re.I):
        return 'divisor'
    return 'conteudo'


# ---------- textos ----------
def trocar_texto(shape, linhas):
    """Escreve `linhas` (uma por parágrafo) reaproveitando o formato do 1º trecho de cada parágrafo."""
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    for k, txt in enumerate(linhas):
        p = paras[min(k, len(paras) - 1)]
        if k >= len(paras):
            novo = copy.deepcopy(p._p); tf._txBody.append(novo)
            p = tf.paragraphs[-1]
        runs = p.runs
        if runs:
            runs[0].text = txt
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run().text = txt
    for p in list(tf.paragraphs)[len(linhas):]:
        p._p.getparent().remove(p._p)


def titulo_com_negrito(shape, runs_antigos):
    """Título novo com os trechos do antigo; mantém negrito onde havia, tira realce/cores antigas."""
    p = shape.text_frame.paragraphs[0]
    modelo = copy.deepcopy(p.runs[0]._r) if p.runs else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for k, (txt, bold) in enumerate(runs_antigos):
        if modelo is None:
            r = p.add_run(); r.text = txt
        else:
            el = copy.deepcopy(modelo); el.find(qa('t')).text = txt
            p._p.insert(len(p._p) - (1 if p._p[-1].tag == qa('endParaRPr') else 0), el)
            r = p.runs[-1]
        if bold:
            r.font.bold = True
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def runs_do_titulo(sh):
    out = []
    for k, p in enumerate(sh.text_frame.paragraphs):
        if k and out:
            out.append((' ', False))
        for r in p.runs:
            if r.text:
                out.append((r.text.replace('\x0b', ' '), bool(r.font.bold)))
    # junta trechos vizinhos de mesmo estilo
    junt = []
    for t, b in out:
        if junt and junt[-1][1] == b:
            junt[-1] = (junt[-1][0] + t, b)
        else:
            junt.append((t, b))
    if junt:
        junt[0] = (junt[0][0].lstrip(), junt[0][1]); junt[-1] = (junt[-1][0].rstrip(), junt[-1][1])
    return [(re.sub(r'\s+', ' ', t), b) for t, b in junt if t]


def textos_da_capa(slide):
    linhas = []
    for sh in slide.shapes:
        if sh.has_text_frame and not re.search(r'lorem ipsum', sh.text_frame.text, re.I):
            linhas += [x.strip(' :') for x in re.split(r'[\n\x0b]', sh.text_frame.text) if x.strip()]
    data = next((x for x in linhas if re.search(r'offsite|20\d\d', x, re.I)), '')
    resto = [x for x in linhas if x != data]
    titulo = next((x for x in resto if re.search(r'an[aá]lise', x, re.I)), resto[0] if resto else 'TÍTULO')
    sub = next((x for x in resto if x != titulo), '')
    return titulo, sub, data


# ---------- montagem ----------
def clonar_modelo(dst, modelo, cop, pular=()):
    s = dst.slides.add_slide(modelo.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in modelo.shapes:
        if eh_thinkcell(sh) or sh.name in pular:
            continue
        cop.copiar(sh._element, modelo.part, s)
    return s


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--deck', required=True); ap.add_argument('--saida', required=True)
    ap.add_argument('--template', default=PADRAO); ap.add_argument('--etiqueta', help='texto da etiqueta cinza (padrão: título da capa + subtítulo)')
    ap.add_argument('--data', help='texto da data na capa (padrão: o da capa antiga)'); ap.add_argument('--conferencia')
    args = ap.parse_args()

    src = Presentation(args.deck)
    dst = abrir_template(args.template)
    if (src.slide_width, src.slide_height) != (dst.slide_width, dst.slide_height):
        sys.exit('O deck e o template têm tamanhos de slide diferentes; ajuste antes de aplicar.')
    mod = list(dst.slides)
    por_layout = lambda nome, cond=lambda s: True: next((s for s in mod if s.slide_layout.name == nome and cond(s)), None)
    m_capa = mod[0]
    # slide de conteúdo do template: etiqueta "1) xx: xx" e título "XX", sem outras caixas de texto
    m_cont = por_layout('Corpo_texto pequeno', lambda s: any(sh.has_text_frame and re.match(r'\s*[1-9]\)', sh.text_frame.text) for sh in s.shapes)
                        and not any(sh.name.startswith('TextBox') for sh in s.shapes)) or por_layout('Corpo_texto pequeno')
    m_agenda = por_layout('Texto e Título Vertical')
    m_fim = mod[-1]
    if not (m_cont and m_agenda):
        sys.exit('Não reconheci os slides-modelo do template (conteúdo "Corpo_texto pequeno" e agenda "Texto e Título Vertical").')

    cop = Copiador(dst)
    tit_capa, sub_capa, data_capa = textos_da_capa(src.slides[0])
    etiqueta_base = args.etiqueta or (f'{tit_capa.rstrip(":")}: {sub_capa}' if sub_capa else tit_capa)
    secoes = [titulo_antigo(s).text_frame.text.strip() if titulo_antigo(s) else f'Seção {k + 1}'
              for k, s in enumerate(x for i, x in enumerate(src.slides) if tipo_slide(i, x) == 'divisor')]
    mapa, conf, secao_atual = {}, [], None

    for i, old in enumerate(src.slides):
        tipo = tipo_slide(i, old)
        if tipo == 'capa':
            s = clonar_modelo(dst, m_capa, cop)
            for sh in s.shapes:
                if not sh.has_text_frame:
                    continue
                if re.search(r'lorem ipsum', sh.text_frame.text, re.I):
                    sh._element.getparent().remove(sh._element)
                elif 'TÍTULO' in sh.text_frame.text:
                    ps = sh.text_frame.paragraphs
                    ps[0].runs[0].text = tit_capa.rstrip(':').upper()
                    ps[1].runs[0].text = sub_capa
                    ultimo = [p for p in ps if p.runs][-1]
                    ultimo.runs[-1].text = f'| {args.data or data_capa or "mês/ano"} |'
            conf.append(f'- Slide {i + 1} → capa do template: "{tit_capa}" / "{sub_capa}" / "{args.data or data_capa}"')
        elif tipo == 'encerramento':
            s = clonar_modelo(dst, m_fim, cop)
            conf.append(f'- Slide {i + 1} → encerramento do template')
        elif tipo == 'divisor':
            t = titulo_antigo(old); nome = t.text_frame.text.strip() if t else ''
            secao_atual = secoes.index(nome) + 1 if nome in secoes else None
            s = clonar_modelo(dst, m_agenda, cop)
            itens = [(a, b) for a, b in zip([x for x in s.shapes if x.name.startswith('CaixaDeTexto')],
                                            [x for x in s.shapes if x.name.startswith('Rectangle')])]
            passo = (itens[1][0].top - itens[0][0].top) if len(itens) > 1 else Inches(1.38)
            caixa0, num0 = itens[0]
            for _, (cx, nr) in enumerate(itens):
                cx._element.getparent().remove(cx._element); nr._element.getparent().remove(nr._element)
            for k, nome_k in enumerate(secoes):
                cx = copy.deepcopy(caixa0._element); nr = copy.deepcopy(num0._element)
                s.shapes._spTree.append(cx); s.shapes._spTree.append(nr)
                cxs, nrs = s.shapes[-2], s.shapes[-1]
                cxs.top = caixa0.top + passo * k; nrs.top = num0.top + passo * k
                trocar_texto(cxs, [nome_k]); trocar_texto(nrs, [str(k + 1)])
                if k + 1 != secao_atual:  # seção que não é a atual fica em cinza, como no template
                    for el in list(cxs._element.iter(qa('solidFill'))) + [nrs._element.find(qp('spPr')).find(qa('solidFill'))]:
                        if el is None:
                            continue
                        for c in list(el):
                            el.remove(c)
                        sc = etree.SubElement(el, qa('schemeClr'), val='tx1')
                        etree.SubElement(sc, qa('lumMod'), val='65000'); etree.SubElement(sc, qa('lumOff'), val='35000')
            conf.append(f'- Slide {i + 1} → agenda do template, seção {secao_atual}: "{nome}"')
        else:
            t = titulo_antigo(old)
            s = clonar_modelo(dst, m_cont, cop)
            textos = [sh for sh in s.shapes if sh.has_text_frame and sh.name.startswith('Text Placeholder')]
            tag = min(textos, key=lambda x: x.top); tit = max(textos, key=lambda x: x.top)
            n = f'{secao_atual}) ' if secao_atual else ''
            trocar_texto(tag, [n + (secoes[secao_atual - 1] if secao_atual else etiqueta_base)])
            if t is not None:
                titulo_com_negrito(tit, runs_do_titulo(t))
            else:
                trocar_texto(tit, ['']); conf.append(f'- Slide {i + 1}: sem título no original; o título do template ficou vazio.')
            invade = []
            for sh in old.shapes:
                if sh is t or (t is not None and sh.shape_id == t.shape_id) or eh_thinkcell(sh):
                    continue
                el = cop.copiar(sh._element, old.part, s)
                sem_placeholder(el, sh)
                if sh.top is not None and sh.top < TOPO_CONTEUDO and (sh.top + (sh.height or 0)) > Inches(0.35) and sh.width and sh.width > Inches(0.3):
                    invade.append(sh.name)
            # notas do apresentador
            if old.has_notes_slide and old.notes_slide.notes_text_frame.text.strip():
                s.notes_slide.notes_text_frame.text = old.notes_slide.notes_text_frame.text
            titulo_txt = t.text_frame.text.strip().replace('\n', ' ') if t is not None else ''
            conf.append(f'- Slide {i + 1} → conteúdo: "{titulo_txt[:90]}"' +
                        (f' · **confira**: {len(invade)} forma(s) na faixa do título ({", ".join(invade[:4])})' if invade else ''))
        mapa[id(old.part)] = s

    cop.resolver_links(mapa)
    # tira os slides de exemplo do template
    ids = dst.slides._sldIdLst
    for s in mod:
        for sid in list(ids):
            if dst.part.related_part(sid.rId) is s.part:
                dst.part.drop_rel(sid.rId); ids.remove(sid)
    dst.save(args.saida)
    path = args.conferencia or os.path.splitext(args.saida)[0] + '_template.md'
    open(path, 'w', encoding='utf-8').write(
        f'# Deck no template: {os.path.basename(args.saida)}\n\n- Original: `{os.path.basename(args.deck)}`\n'
        f'- Template: `{os.path.basename(args.template)}`\n- Etiqueta: "{etiqueta_base}"\n\n## Slides\n' + '\n'.join(conf) + '\n')
    print(f'{args.saida}\n{path}\n{len(src.slides)} slides')


if __name__ == '__main__':
    main()

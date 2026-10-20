"""Atualiza o PowerPoint "Radar Estratégia – Modelo 2" com os dados do painel e a leitura do time.

O que faz, sem quebrar o vínculo dos gráficos com a Grande Base no SharePoint:
  1. Dados: para cada série de cada gráfico, lê as células da Grande Base usando as mesmas faixas
     do painel (pipeline/radar/faixas.py); se a faixa do painel cresceu, estende a faixa vinculada
     do gráfico (assim o "Atualizar links" do PowerPoint também passa a ler o mês novo). Os valores
     ficam gravados no gráfico, então o deck abre atualizado mesmo sem acesso ao SharePoint.
  2. Rótulos do último ponto: move para o último dado novo.
  3. "Dados atualizados até mmm/aa": troca pela data do último dado observado de cada gráfico.
  4. Textos: título, texto de apoio e caixas de leitura de cada slide a partir de
     pipeline/radar/leitura_time.json (mapa de caixas em deck_map_modelo2.json), mantendo a
     formatação do slide (títulos em caixa alta, destaques em negrito/laranja).
  5. Se a leitura não estiver revisada, marca os slides com "RASCUNHO".
  6. Gera uma conferência em Markdown: o que mudou, revisões de dados na fonte e o que revisar à mão.

Uso:
  python atualizar_deck.py --projeto <painel-estrategico> --grande-base "<Grande Base>.xlsx" \
      --deck "<Radar ... Modelo 2>.pptx" --saida "Radar_Modelo2_atualizado.pptx" \
      [--leitura <leitura_time.json>] [--capa "novembro, 2026"] [--conferencia conferencia_deck.md]
  python atualizar_deck.py --deck "<deck>.pptx" --listar      # lista slides, caixas de texto e gráficos
"""
import argparse, copy, datetime, importlib.util, json, os, re, sys, warnings
warnings.filterwarnings('ignore')
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
q = lambda t: f'{{{C}}}{t}'
qa = lambda t: f'{{{A}}}{t}'
MES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
REF = re.compile(r"^'?(?P<sheet>[^'!]+)'?!\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$")
TOKEN = re.compile(r'\b(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)/\d{2}\b')
HOJE = datetime.date.today()


# ---------- utilidades ----------
def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def serial(v):
    if isinstance(v, datetime.datetime):
        v = v.date()
    if isinstance(v, datetime.date):
        return float((v - datetime.date(1899, 12, 30)).days)
    return num(v)


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # número de série do Excel
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    return None


def mmmaa(d):
    return f'{MES[d.month - 1]}/{str(d.year)[2:]}'


def read_cache(cache):
    if cache is None:
        return []
    n = int(cache.find(q('ptCount')).get('val'))
    out = [None] * n
    for pt in cache.findall(q('pt')):
        i = int(pt.get('idx'))
        if i < n:
            t = pt.find(q('v')).text
            try:
                out[i] = float(t)
            except (TypeError, ValueError):
                out[i] = t
    return out


def write_cache(cache, vals, text=False):
    for pt in cache.findall(q('pt')):
        cache.remove(pt)
    cache.find(q('ptCount')).set('val', str(len(vals)))
    anchor = cache.find(q('extLst'))
    for i, v in enumerate(vals):
        if v is None:
            continue
        pt = etree.Element(q('pt'), idx=str(i))
        etree.SubElement(pt, q('v')).text = (str(v) if text else repr(float(v)))
        if anchor is not None:
            anchor.addprevious(pt)
        else:
            cache.append(pt)


def load_specs(projeto):
    path = os.path.join(projeto, 'pipeline', 'radar', 'faixas.py')
    if not os.path.exists(path):
        sys.exit(f'Não encontrei {path}: aponte --projeto para a pasta do painel-estrategico.')
    spec = importlib.util.spec_from_file_location('faixas', path)
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    look = {}
    for chave, series in mod.SPECS.items():
        for s in series:
            look.setdefault((s['sheet'], s['col']), []).append((s['r1'], s['r2'], chave, s['label']))
    return look


def cell(wb, sheet, col, r):
    return wb[sheet][f'{col}{r}'].value


# ---------- 1–2. dados dos gráficos ----------
def atualizar_serie(ser, wb, look, ctx, log):
    val = ser.find(q('val'))
    vf = val.find('.//' + q('f')) if val is not None else None
    m = REF.match(vf.text or '') if vf is not None else None
    item = dict(ctx, serie=None, faixa_antes=vf.text if vf is not None else '', faixa_depois='', ultimo_antes=None, ultimo_depois=None,
                revisoes=0, rotulo_movido=False, obs='')
    tx = ser.find(q('tx'))
    old_name = read_cache(tx.find('.//' + q('strCache'))) if tx is not None else []
    item['serie'] = (old_name[0] if old_name else '') or f"série {ctx['i'] + 1}"
    if not m or not m.group('r2'):
        item['obs'] = 'dados fixos no gráfico (sem vínculo com a Grande Base): mantidos'
        log.append(item); return None
    sh, col, a, b = m.group('sheet'), m.group('c1'), int(m.group('r1')), int(m.group('r2'))
    if sh not in wb.sheetnames:
        item['obs'] = f'aba "{sh}" não existe na Grande Base: dados mantidos como estavam no deck'
        log.append(item); return None
    cand = look.get((sh, col), [])
    spec = next((c for c in cand if c[0] == a), None) or (max(cand, key=lambda c: min(c[1], b) - max(c[0], a)) if cand else None)
    nb = spec[1] if spec else b
    delta = nb - b
    if not spec:
        item['obs'] = 'série não está no painel: faixa mantida, valores relidos da Grande Base'
    vals = [num(cell(wb, sh, col, r)) for r in range(a, nb + 1)]
    vcache = val.find('.//' + q('numCache'))
    old_vals = read_cache(vcache)
    item['revisoes'] = sum(1 for o, n in zip(old_vals, vals) if isinstance(o, float) and n is not None and abs(o - n) > 1e-9 * max(1, abs(o)))
    write_cache(vcache, vals)
    vf.text = f"{vf.text.split('!')[0]}!${col}${a}:${col}${nb}"
    item['faixa_depois'] = vf.text
    # categorias
    cats = []
    cat = ser.find(q('cat'))
    def cat_cache(c):
        n = c.find('.//' + q('numCache'))
        return n if n is not None else c.find('.//' + q('strCache'))
    old_cats = read_cache(cat_cache(cat)) if cat is not None else []
    if cat is not None:
        cf = cat.find('.//' + q('f'))
        cm = REF.match(cf.text or '') if cf is not None else None
        if cm and cm.group('r2') and cm.group('sheet') in wb.sheetnames:
            cs, cc, ca, cb = cm.group('sheet'), cm.group('c1'), int(cm.group('r1')), int(cm.group('r2'))
            raw = [cell(wb, cs, cc, r) for r in range(ca, cb + delta + 1)]
            cf.text = f"{cf.text.split('!')[0]}!${cc}${ca}:${cc}${cb + delta}"
            nc, sc = cat.find('.//' + q('numCache')), cat.find('.//' + q('strCache'))
            if nc is not None:
                write_cache(nc, [serial(v) for v in raw])
            elif sc is not None:
                write_cache(sc, [(str(int(v)) if isinstance(v, float) and v.is_integer() else (str(v).strip() if v is not None else None)) for v in raw], text=True)
            cats = raw
        else:
            cats = old_cats
    # nome da série
    if tx is not None:
        tf = tx.find('.//' + q('f')); tm = REF.match(tf.text or '') if tf is not None else None
        if tm and tm.group('sheet') in wb.sheetnames:
            nm = cell(wb, tm.group('sheet'), tm.group('c1'), int(tm.group('r1')))
            if nm:
                write_cache(tx.find('.//' + q('strCache')), [str(nm)], text=True); item['serie'] = str(nm)

    # último ponto observado (até hoje; projeções ficam de fora)
    def data_do_ponto(cs, i):
        if i < len(cs):
            return as_date(cs[i])
        # mais valores que datas na faixa do gráfico: estende a série mensal a partir da última data
        ds = [as_date(c) for c in cs]
        if len(ds) >= 2 and ds[-1] and ds[-2] and 27 <= (ds[-1] - ds[-2]).days <= 32:
            k = i - (len(cs) - 1); y, m = ds[-1].year, ds[-1].month + k
            y += (m - 1) // 12; m = (m - 1) % 12 + 1
            return datetime.date(y, m, 1)
        return None

    def ultimo(vs, cs):
        for i in range(len(vs) - 1, -1, -1):
            if vs[i] is None:
                continue
            d = data_do_ponto(cs, i)
            if d and d > HOJE:
                continue
            return i, d
        return None, None
    oi, od = ultimo(old_vals, old_cats)
    ni, nd = ultimo(vals, cats)
    li_old = max((i for i, v in enumerate(old_vals) if v is not None), default=None)
    li_new = max((i for i, v in enumerate(vals) if v is not None), default=None)
    item['ultimo_antes'] = mmmaa(od) if od else (f'ponto {oi + 1}' if oi is not None else None)
    item['ultimo_depois'] = mmmaa(nd) if nd else (f'ponto {ni + 1}' if ni is not None else None)
    # rótulo do último ponto acompanha o dado novo
    dl = ser.find(q('dLbls'))
    if dl is not None and li_old is not None and li_new is not None and li_new != li_old:
        for d in dl.findall(q('dLbl')):
            ix = d.find(q('idx'))
            if ix is not None and int(ix.get('val')) == li_old:
                ix.set('val', str(li_new)); item['rotulo_movido'] = True
    log.append(item)
    return (od, nd)


def atualizar_graficos(prs, wb, look):
    log, datas = [], {}
    for sn, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not getattr(shp, 'has_chart', False):
                continue
            for i, ser in enumerate(shp.chart._chartSpace.iter(q('ser'))):
                d = atualizar_serie(ser, wb, look, dict(slide=sn, grafico=shp.name, i=i), log)
                datas.setdefault((sn, shp.shape_id), []).append(d)
    return log, datas


# ---------- 3. "Dados atualizados até" ----------
def atualizar_datas(prs, datas, pend):
    feitos = []
    for sn, slide in enumerate(prs.slides, 1):
        charts = [s for s in slide.shapes if getattr(s, 'has_chart', False)]
        for shp in slide.shapes:
            if not shp.has_text_frame or 'Dados atualizados' not in shp.text_frame.text:
                continue
            cx = shp.left + shp.width / 2
            acima = [c for c in charts if c.top < shp.top] or charts
            if not acima:
                continue
            ch = min(acima, key=lambda c: (abs(c.left + c.width / 2 - cx), shp.top - c.top))
            pares = [x for x in datas.get((sn, ch.shape_id), []) if x and x[1]]
            antes = shp.text_frame.text
            toks = [m.group(0) for m in TOKEN.finditer(antes)]
            if not pares or not toks:
                continue
            if len(toks) == 1:
                novos = [mmmaa(max(n for _, n in pares))]
            else:
                # cada data do texto corresponde à série cujo último dado (no deck antigo) era essa data
                novos = [None] * len(toks); livres = list(range(len(pares)))
                for k, tk in enumerate(toks):
                    m_ = [j for j in livres if pares[j][0] and mmmaa(pares[j][0]) == tk]
                    if len(m_) == 1:
                        novos[k] = mmmaa(pares[m_[0]][1]); livres.remove(m_[0])
                faltam = [k for k, v in enumerate(novos) if v is None]
                if len(faltam) == 1 and len(livres) == 1:
                    j = livres[0]; k = faltam[0]
                    pend.append(f'Slide {sn}: o texto anterior dizia "{toks[k]}", mas a série já ia até {mmmaa(pares[j][0]) if pares[j][0] else "?"} no deck antigo; corrigido para {mmmaa(pares[j][1])}. Confira.')
                    novos[k] = mmmaa(pares[j][1])
                if None in novos:
                    novos = None
                if novos is None:
                    pend.append(f'Slide {sn}: não consegui associar as datas de "{antes.strip()}" às séries; confira à mão.')
                    continue
            fila = list(novos)
            for p in shp.text_frame.paragraphs:
                runs = p.runs
                full = ''.join(r.text for r in runs)
                if not TOKEN.search(full):
                    continue
                novo = TOKEN.sub(lambda m: fila.pop(0) if fila else m.group(0), full)
                if len(novo) != len(full):  # "mmm/aa" sempre tem 6 caracteres; por segurança
                    pend.append(f'Slide {sn}: não consegui trocar a data em "{antes.strip()}"; confira à mão.'); continue
                pos = 0
                for r in runs:  # redistribui o texto pelos mesmos trechos, preservando a formatação
                    n = len(r.text); r.text = novo[pos:pos + n]; pos += n
            if shp.text_frame.text != antes:
                feitos.append((sn, antes.strip(), shp.text_frame.text.strip()))
    return feitos


# ---------- 4–5. textos ----------
def set_rich(shape, texto, upper=False, laranja='FFA74F', minimo=9, min_linhas=1):
    """Troca o texto mantendo a formatação do slide.

    `texto` pode ser uma frase ou uma lista (um item por bullet/parágrafo). Cada parágrafo novo copia o
    formato do parágrafo original na mesma posição (marcador, recuo, alinhamento); se houver mais itens
    que parágrafos, repete o formato do último. Destaques entre ** usam o formato dos trechos em negrito
    do original (ou negrito + laranja). Se o texto novo for mais longo que o antigo, reduz a fonte na
    proporção (até `minimo` pt). Devolve o tamanho final da fonte, se reduziu.
    """
    itens = [x for x in (texto if isinstance(texto, list) else [texto]) if x]
    tf = shape.text_frame
    paras = tf.paragraphs
    len_antigo = max(1, len(tf.text.strip()))
    len_novo = max(1, sum(len(x.replace('**', '').strip()) for x in itens))
    runs = [r for p in paras for r in p.runs]
    # formato "normal" e "destaque" = os trechos mais longos de cada tipo no original (evita copiar o
    # formato de um rótulo curto, como "Protagonismo geopolítico:" sublinhado)
    normais = [r for r in runs if not r.font.bold]; negritos = [r for r in runs if r.font.bold]
    base = max(normais, key=lambda r: len(r.text)) if normais else (runs[0] if runs else None)
    dest = max(negritos, key=lambda r: len(r.text)) if negritos else None
    nr = copy.deepcopy(base._r.find(qa('rPr'))) if base is not None and base._r.find(qa('rPr')) is not None else etree.Element(qa('rPr'))
    if dest is not None and dest._r.find(qa('rPr')) is not None:
        hr = copy.deepcopy(dest._r.find(qa('rPr')))
    else:
        hr = copy.deepcopy(nr); hr.set('b', '1')
        for f in hr.findall(qa('solidFill')):
            hr.remove(f)
        sf = etree.Element(qa('solidFill')); etree.SubElement(sf, qa('srgbClr'), val=laranja)
        hr.insert(0, sf)
    final = None
    sz0 = nr.get('sz') or (str(int(base.font.size.pt * 100)) if base is not None and base.font.size else None)
    if sz0:
        # quantos caracteres cabem na caixa em cada tamanho de fonte (estimativa: largura média do caractere
        # ~0,55 em minúsculas e ~0,65 em caixa alta; altura de linha 1,2x a fonte; ao menos o texto antigo cabe)
        larg = shape.width / 12700 - 14; alt = shape.height / 12700 - 7
        def cabe(pt):
            linhas = max(min_linhas, int(alt // (pt * 1.2)))
            por_linha = larg / (pt * (0.65 if upper else 0.55))
            return max(len_antigo, por_linha * linhas * 0.92) - 3 * (len(itens) - 1)
        pt = int(sz0) / 100
        while pt > minimo and len_novo > cabe(pt):
            pt -= 0.5
        if pt < int(sz0) / 100:
            for rp in (nr, hr):
                rp.set('sz', str(int(pt * 100)))
            final = pt
    modelos = [copy.deepcopy(p._p) for p in paras]
    body = tf._txBody
    for p in paras:
        body.remove(p._p)
    for k, item in enumerate(itens):
        p0 = copy.deepcopy(modelos[min(k, len(modelos) - 1)])
        for el in list(p0):
            if el.tag in (qa('r'), qa('br'), qa('fld')):
                p0.remove(el)
        end = p0.find(qa('endParaRPr'))
        for i, seg in enumerate(item.split('**')):
            if not seg:
                continue
            r = etree.Element(qa('r')); r.append(copy.deepcopy(hr if i % 2 else nr))
            etree.SubElement(r, qa('t')).text = seg.upper() if upper else seg
            if end is not None:
                end.addprevious(r)
            else:
                p0.append(r)
        body.append(p0)
    return final


def carimbo(slide, texto):
    tb = slide.shapes.add_textbox(Inches(10.75), Inches(0.0), Inches(2.58), Inches(0.17))
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
    tb.fill.solid(); tb.fill.fore_color.rgb = RGBColor(0xB3, 0x26, 0x1E)
    p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = texto
    r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.name = 'Trebuchet MS'
    tb.name = 'Carimbo rascunho'


def atualizar_textos(prs, leitura, mapa, pend):
    feitos = []
    rasc = leitura.get('status') != 'revisado'
    top = leitura.get('topicos', {}); extra = leitura.get('deck_modelo2', {})
    for sn_txt, cfg in mapa['slides'].items():
        sn = int(sn_txt)
        if sn > len(prs.slides):
            continue
        slide = prs.slides[sn - 1]; t = top.get(cfg['topico'], {}); ex = extra.get(sn_txt, {})
        shapes = {s.shape_id: s for s in slide.shapes}
        mudou = False
        for papel in ('titulo', 'subtitulo', 'caixa_esq', 'caixa_dir'):
            if papel not in cfg:
                continue
            shp = shapes.get(cfg[papel])
            if shp is None or not shp.has_text_frame:
                pend.append(f'Slide {sn}: caixa "{papel}" (id {cfg[papel]}) não encontrada; atualize deck_map_modelo2.json (rode com --listar).'); continue
            texto = ex.get(papel) or ({'titulo': t.get('titulo'), 'subtitulo': t.get('texto')}.get(papel))
            if not texto:
                pend.append(f'Slide {sn}: sem texto para "{papel}" na leitura; mantido o da edição anterior.'); continue
            antigo = shp.text_frame.text
            upper = papel == 'titulo' and antigo.strip() == antigo.strip().upper()
            if papel == 'titulo':
                # título justificado espalha as palavras quando quebra em 2 linhas: alinha à esquerda
                for para in shp.text_frame.paragraphs:
                    pPr = para._p.find(qa('pPr'))
                    if pPr is not None and pPr.get('algn') == 'just':
                        pPr.set('algn', 'l')
                # termina o título antes de elementos à direita na mesma faixa (ex.: marcador verde do modelo)
                dir_ = [o for o in slide.shapes if o.shape_id != shp.shape_id and o.width < Inches(2) and o.height < Inches(1.5)
                        and o.left > shp.left + shp.width * 0.6
                        and o.left < shp.left + shp.width and o.top < shp.top + shp.height and o.top + o.height > shp.top]
                if dir_:
                    shp.width = min(o.left for o in dir_) - shp.left - 45720
            if papel == 'subtitulo' and getattr(slide, '_desce_sub', 0):
                d = slide._desce_sub  # título ocupou 2 linhas: desce os bullets e mantém a base da caixa
                shp.top = shp.top + d; shp.height = max(Inches(0.5), shp.height - d)
            sz = set_rich(shp, texto, upper=upper, laranja='FFA74F' if papel.startswith('caixa') else 'ED7D31', minimo=14 if papel == 'titulo' else 9, min_linhas=1 if papel == 'subtitulo' else 2)
            feitos.append((sn, papel, antigo.replace('\n', ' / ')[:90], (' / '.join(texto) if isinstance(texto, list) else texto).replace('**', '')[:90]))
            if papel == 'titulo' and 'subtitulo' in cfg:
                sub = shapes.get(cfg['subtitulo'])
                pt = sz or (shp.text_frame.paragraphs[0].runs[0].font.size.pt if shp.text_frame.paragraphs[0].runs and shp.text_frame.paragraphs[0].runs[0].font.size else 20)
                por_linha = (shp.width / 12700 - 14) / (pt * (0.65 if upper else 0.55)) * 0.92
                n_lin = -(-len(texto.replace('**', '')) // int(por_linha))
                folga = (sub.top - (shp.top + shp.height)) if sub is not None else 0
                if sub is not None and n_lin >= 2 and folga < Inches(0.25):
                    slide._desce_sub = int(Inches(pt * 1.2 / 72 * (n_lin - 1)) - max(0, folga) + Inches(0.05))
            if sz:
                pend.append(f'Slide {sn}: {papel} mais longo que o da edição anterior; fonte reduzida para {sz:g} pt. Se ficar pequeno, encurte o texto na leitura.')
            mudou = True
        if mudou and rasc:
            carimbo(slide, 'RASCUNHO · texto aguardando revisão do time')
    for sn_txt, motivo in mapa.get('manter', {}).items():
        pend.append(f'Slide {sn_txt}: {motivo}.')
    return feitos, rasc


def rotulos_manuais(prs, pend):
    """Caixas de texto com números digitados à mão em cima dos gráficos (ex.: '5,9%')."""
    for sn, slide in enumerate(prs.slides, 1):
        charts = [s for s in slide.shapes if getattr(s, 'has_chart', False)]
        for shp in slide.shapes:
            if not shp.has_text_frame or getattr(shp, 'has_chart', False):
                continue
            tx = shp.text_frame.text.strip()
            if re.fullmatch(r'[-−]?\d{1,3}([.,]\d+)?\s*%?', tx.split('\n')[0].strip()) and any(
                    c.left <= shp.left <= c.left + c.width and c.top <= shp.top <= c.top + c.height for c in charts):
                pend.append(f'Slide {sn}: rótulo digitado à mão "{tx.splitlines()[0]}" sobre o gráfico; confira com o último valor da série.')


# ---------- conferência ----------
def conferencia(path, log, datas_feitas, textos, pend, rasc, args):
    L = [f'# Conferência do deck Radar Modelo 2 ({HOJE.strftime("%d/%m/%Y")})', '',
         f'- Base: `{os.path.basename(args.deck)}` · Grande Base: `{os.path.basename(args.grande_base)}`',
         f'- Leitura: `{os.path.basename(args.leitura)}` · status **{"RASCUNHO (slides com carimbo)" if rasc else "revisada"}**', '']
    ext = [x for x in log if x['faixa_depois'] and x['faixa_antes'] != x['faixa_depois']]
    rev = [x for x in log if x['revisoes']]
    L += ['## Resumo', f'- Séries atualizadas: {sum(1 for x in log if x["faixa_depois"])} de {len(log)}',
          f'- Séries com faixa estendida (mês novo): {len(ext)}', f'- Séries com dados revisados na fonte: {len(rev)}',
          f'- Rótulos de último ponto movidos: {sum(1 for x in log if x["rotulo_movido"])}',
          f'- Datas "Dados atualizados até" trocadas: {len(datas_feitas)} · textos trocados: {len(textos)}', '']
    L += ['## Revisar à mão'] + [f'- {p}' for p in pend] + ([f'- {x["slide"]} · {x["grafico"]} · {x["serie"]}: {x["obs"]}' for x in log if x['obs']]) + ['']
    L += ['## Gráficos', '| Slide | Gráfico | Série | Faixa | Último dado | Revisões | Rótulo movido |', '|---|---|---|---|---|---|---|']
    for x in log:
        fx = x['faixa_depois'] if x['faixa_antes'] == x['faixa_depois'] else f"{x['faixa_antes']} → {x['faixa_depois']}"
        ud = x['ultimo_depois'] if x['ultimo_antes'] == x['ultimo_depois'] else f"{x['ultimo_antes']} → {x['ultimo_depois']}"
        L.append(f"| {x['slide']} | {x['grafico']} | {x['serie']} | {fx or '–'} | {ud or '–'} | {x['revisoes'] or ''} | {'sim' if x['rotulo_movido'] else ''} |")
    L += ['', '## Datas', ''] + [f'- Slide {s}: "{a}" → "{b}"' for s, a, b in datas_feitas]
    L += ['', '## Textos', ''] + [f'- Slide {s} · {p}: "{a}…" → "{b}…"' for s, p, a, b in textos]
    open(path, 'w', encoding='utf-8').write('\n'.join(L) + '\n')


def listar(deck):
    prs = Presentation(deck)
    for sn, slide in enumerate(prs.slides, 1):
        print(f'== slide {sn}')
        for s in slide.shapes:
            kind = 'GRÁFICO' if getattr(s, 'has_chart', False) else ('texto' if s.has_text_frame and s.text_frame.text.strip() else None)
            if kind:
                print(f'  id={s.shape_id:<4} {kind:8} {s.name:<22} x={s.left / 914400:.2f} y={s.top / 914400:.2f} | {(s.text_frame.text if kind == "texto" else "")[:90]!r}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--deck', required=True); ap.add_argument('--listar', action='store_true')
    ap.add_argument('--projeto', default='.'); ap.add_argument('--grande-base'); ap.add_argument('--saida')
    ap.add_argument('--leitura'); ap.add_argument('--mapa', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'deck_map_modelo2.json'))
    ap.add_argument('--capa', help='texto da data na capa, ex.: "novembro, 2026"'); ap.add_argument('--conferencia')
    ap.add_argument('--sem-textos', action='store_true', help='só atualiza os dados')
    args = ap.parse_args()
    if args.listar:
        return listar(args.deck)
    if not (args.grande_base and args.saida):
        sys.exit('Informe --grande-base e --saida.')
    args.leitura = args.leitura or os.path.join(args.projeto, 'pipeline', 'radar', 'leitura_time.json')
    look = load_specs(args.projeto)
    wb = load_workbook(args.grande_base, data_only=True)
    prs = Presentation(args.deck)
    pend = []
    log, datas = atualizar_graficos(prs, wb, look)
    datas_feitas = atualizar_datas(prs, datas, pend)
    textos, rasc = [], False
    if not args.sem_textos:
        if os.path.exists(args.leitura):
            textos, rasc = atualizar_textos(prs, json.load(open(args.leitura, encoding='utf-8')), json.load(open(args.mapa, encoding='utf-8')), pend)
        else:
            pend.append(f'Leitura {args.leitura} não encontrada: textos mantidos da edição anterior.')
    if args.capa:
        s1 = prs.slides[0]
        alvo = next((s for s in s1.shapes if s.has_text_frame and re.search(r'\b20\d\d\b', s.text_frame.text)), None)
        if alvo is not None:
            r = alvo.text_frame.paragraphs[0].runs
            if r:
                r[0].text = args.capa
                for x in r[1:]:
                    x.text = ''
    rotulos_manuais(prs, pend)
    prs.save(args.saida)
    conf = args.conferencia or os.path.splitext(args.saida)[0] + '_conferencia.md'
    conferencia(conf, log, datas_feitas, textos, pend, rasc, args)
    print(json.dumps({'saida': os.path.abspath(args.saida), 'conferencia': os.path.abspath(conf), 'series': len(log),
                      'series_atualizadas': sum(1 for x in log if x['faixa_depois']),
                      'faixas_estendidas': sum(1 for x in log if x['faixa_depois'] and x['faixa_antes'] != x['faixa_depois']),
                      'series_com_revisao': sum(1 for x in log if x['revisoes']), 'rotulos_movidos': sum(1 for x in log if x['rotulo_movido']),
                      'datas_trocadas': len(datas_feitas), 'textos_trocados': len(textos), 'rascunho': rasc, 'revisar_a_mao': len(pend)},
                     ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()

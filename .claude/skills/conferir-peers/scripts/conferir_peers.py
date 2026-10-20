"""Conferência dos números da Análise de Peers.

Lê a planilha Análise de Competidores com os mesmos scripts do painel (pasta pipeline/peers do
projeto painel-estrategico), confere cada valor com a célula de origem e gera uma planilha Excel com:
  - Leia-me: como usar, critérios e quadro de andamento (fórmulas);
  - Pontos de atenção: alertas por prioridade, com a célula de origem e colunas para o time preencher;
  - Mudanças (com --anterior): o que mudou em relação à versão anterior da planilha;
  - Base completa: todos os valores do painel com a célula de origem e o tratamento;
  - Em aberto: pontos conhecidos que dependem de decisão (arquivo em_aberto.json, opcional).

Uso:
  python conferir_peers.py --projeto <pasta do painel-estrategico> --planilha "<Análise de Competidores>.xlsx"
         [--anterior "<versão anterior>.xlsx"] [--em-aberto em_aberto.json] [--saida Conferencia_Peers.xlsx]

Imprime um resumo em JSON (contagens por prioridade e tipo de alerta, divergências de leitura).
"""
import argparse, collections, json, os, re, subprocess, sys, tempfile, warnings
warnings.filterwarnings('ignore')
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SHEETS = ['Gráficos - Louças', 'Gráficos - Metais', 'Gráficos Madeira 1', 'Gráficos - Revestimento']
SEG = {'loucas': 'Louças', 'metais': 'Metais', 'madeira': 'Madeira', 'revest': 'Revestimento'}
CAIXA = {'varwk', 'capexe', 'outras', 'fcl', 'varwkr', 'outrasr', 'fclr'}
ERRTXT = {'#N/A': 'N/D (não disponível)', '#DIV/0!': 'divisão por zero', '#VALUE!': 'valor inválido', '#REF!': 'referência inválida'}


# ---------- 1. leitura com os scripts do painel ----------
def pipeline(projeto, planilha):
    """Roda parse_blocks/build/clean do painel numa pasta temporária e devolve peers_clean.json."""
    pp = os.path.abspath(os.path.join(projeto, 'pipeline', 'peers'))
    if not os.path.exists(os.path.join(pp, 'clean.py')):
        sys.exit(f'Não encontrei {pp}/clean.py: aponte --projeto para a pasta do painel-estrategico.')
    w = tempfile.mkdtemp()
    env = {**os.environ, 'PEERS_XLSX': os.path.abspath(planilha)}
    for script, args in (('parse_blocks.py', SHEETS), ('build.py', []), ('clean.py', [])):
        r = subprocess.run([sys.executable, '-I', os.path.join(pp, script), *args], cwd=w, env=env, capture_output=True, text=True)
        if r.returncode != 0:
            sys.exit(f'Falha em {script}:\n{r.stderr[-2000:]}')
    return json.load(open(os.path.join(w, 'peers_clean.json'), encoding='utf-8'))


def ycols(ws, r, c0, years):
    out = {}
    for c in range(c0, c0 + 40):
        v = ws.cell(r, c).value
        if v is None:
            continue
        sv = str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()
        for y in years:
            if sv == y or (y.endswith('LTM') and 'LTM' in sv and sv.split()[0] in y):
                out.setdefault(y, c)
    return out


def linhas(D, planilha):
    """Uma linha por empresa × indicador × ano, com o valor do painel e o da célula de origem."""
    wb = load_workbook(planilha, data_only=True)
    rows, bad = [], []
    for seg, s in D.items():
        for k, m in s['metrics'].items():
            if m.get('calc'):
                for e, ser in m['series'].items():
                    for y, v in zip(m['years'], ser):
                        rows.append(dict(seg=seg, mk=k, ind=m['label'], ent=e, emp=s['ents'][e]['name'], ano=y, fmt=m['fmt'],
                                         painel=v, planilha=None, cel='calculado: receita × margem EBITDA', trat='calculado', why=''))
                continue
            ws = wb[m['sheet']]
            r0 = int(re.search(r'(\d+)$', m['ref']).group(1))
            yc = ycols(ws, r0, ci(m['col']), m['years'])
            for e, ser in m['series'].items():
                ex = (m.get('excl') or {}).get(e, {})
                for y, v in zip(m['years'], ser):
                    cons = m.get('dexCons') and e == 'dexco'
                    r, c = m['rows'].get(e), yc.get(y)
                    cel = f"'{m['sheet']}'!{gl(c)}{r}" if r and c else ''
                    raw = ws.cell(r, c).value if r and c else None
                    trat = 'ok'
                    if y in ex:
                        trat = ex[y]['kind']
                    elif v is None and isinstance(raw, (int, float)) and raw == 0:
                        trat = 'zero = sem dado'
                    elif v is None:
                        trat = 'sem dado'
                    if cons:
                        trat, cel, raw = 'alavancagem consolidada', 'linha Dexco consolidada (Louças/Metais/Revest.)', None
                    rows.append(dict(seg=seg, mk=k, ind=m['label'], ent=e, emp=s['ents'][e]['name'], ano=y, fmt=m['fmt'],
                                     painel=v, planilha=raw, cel=cel, trat=trat, why=ex.get(y, {}).get('why', '')))
                    if not cons and v is not None and not (isinstance(raw, (int, float)) and abs(raw - v) <= 1e-6 * max(1, abs(v))):
                        bad.append(dict(neg=SEG[seg], ind=m['label'], emp=s['ents'][e]['name'], ano=y, painel=v, celula=raw, cel=cel))
    return rows, bad


# ---------- 2. alertas ----------
def fmtv(v, fmt):
    if v is None:
        return '–'
    if fmt == 'pct':
        return f'{v*100:.1f}%'.replace('.', ',')
    if fmt == 'x':
        return f'{v:.1f}x'.replace('.', ',')
    if fmt == 'd':
        return f'{v:.0f} dias'
    if fmt == 'mi':
        return 'R$ ' + f'{v:,.1f}'.replace(',', 'X').replace('.', ',').replace('X', '.') + ' mi'
    return str(v)


def alertas(D, R):
    idx = {(r['seg'], r['mk'], r['ent'], r['ano']): r for r in R}
    val = lambda seg, mk, e, y: (idx.get((seg, mk, e, y)) or {}).get('painel')
    A = []

    def add(pri, r, alerta, det, ref=None):
        A.append(dict(pri=pri, neg=SEG[r['seg']], emp=r['emp'], ind=r['ind'], ano=r['ano'],
                      v=r['painel'] if r['painel'] is not None else r.get('planilha'), fmt=r['fmt'], ref=ref,
                      alerta=alerta, det=det, cel=r['cel'], dex=r['ent'] == 'dexco'))

    for r in R:
        m = D[r['seg']]['metrics'][r['mk']]
        yrs = m['years']; i = yrs.index(r['ano']); v = r['painel']; dex = r['ent'] == 'dexco'
        if r['trat'] in ('erro', 'outlier'):
            add('Alta' if dex else 'Média', r, f"Retirado do painel ({r['trat']})", r['why']); continue
        if v is None:
            continue
        # saltos em relação ao ano anterior (EBITDA calculado da Madeira não entra: repete receita e margem)
        if i > 0 and not m.get('calc') and 'LTM' not in r['ano']:
            p = val(r['seg'], r['mk'], r['ent'], yrs[i - 1])
            if p is not None:
                d = v - p; fmt = r['fmt']; hit = None
                if fmt == 'pct' and r['mk'] not in CAIXA and abs(d) >= .15:
                    hit = 'variação de ' + f'{d*100:+.1f}'.replace('.', ',') + ' p.p.'
                elif fmt == 'pct' and r['mk'] in CAIXA and abs(d) >= 1.0:
                    hit = f'variação de {d*100:+.0f} p.p.'
                elif r['mk'] == 'giro' and abs(d) >= .8:
                    hit = 'variação de ' + f'{d:+.1f}'.replace('.', ',') + 'x'
                elif r['mk'] == 'alav' and abs(d) >= 2.5:
                    hit = 'variação de ' + f'{d:+.1f}'.replace('.', ',') + 'x'
                elif fmt == 'd' and abs(d) >= 60:
                    hit = f'variação de {d:+.0f} dias'
                elif ((r['mk'] in ('rol', 'ci') and p > 0 and abs(d / p) >= .4) or (r['mk'] == 'ebitda' and p >= 50 and abs(d / p) >= .6)) and abs(d) >= 50:
                    hit = f'variação de {d/p*100:+.0f}% no ano'
                if hit:
                    leve = fmt == 'pct' and r['mk'] not in CAIXA and abs(d) < .25
                    pri = 'Alta' if dex else ('Baixa' if leve else 'Média')
                    add(pri, r, 'Salto em relação ao ano anterior', hit + f" (de {fmtv(p, fmt)} para {fmtv(v, fmt)})", p)
                if v == p and v != 0 and r['mk'] != 'alav':
                    add('Alta' if dex else 'Média', r, 'Valor idêntico ao ano anterior', 'pode ser célula copiada ou fórmula puxando o ano errado', p)
        if r['fmt'] == 'mi' and abs(v) >= 500 and abs(v) % 100 == 0:
            add('Alta' if dex else 'Média', r, 'Número redondo', 'valor múltiplo de R$ 100 mi: pode ser estimativa ou dado digitado')
    # consistência entre blocos
    for seg, s in D.items():
        M = s['metrics']
        for e in s['ents']:
            for y in M['rol']['years']:
                rol = val(seg, 'rol', e, y)
                if not rol or rol <= 0:
                    continue
                for mk, num in (('me', 'ebitda'), ('ml', 'll')):
                    if mk in M and num in M and not M[num].get('calc'):
                        a, b = val(seg, mk, e, y), val(seg, num, e, y)
                        if a is not None and b is not None and abs(a - b / rol) > .01:
                            add('Alta' if e == 'dexco' else 'Média', idx[(seg, mk, e, y)], 'Inconsistência entre blocos',
                                f"{M[mk]['label']} {fmtv(a, 'pct')} ≠ {M[num]['label']}/Receita = {fmtv(b / rol, 'pct')}", b / rol)
                if 'giro' in M and 'ci' in M:
                    g, c = val(seg, 'giro', e, y), val(seg, 'ci', e, y)
                    if g and c and c > 0 and abs(g / (rol / c) - 1) > .10:
                        est = s['ents'][e]['country'] != 'BR'
                        add('Baixa' if est else ('Alta' if e == 'dexco' else 'Média'), idx[(seg, 'giro', e, y)], 'Inconsistência entre blocos',
                            f"Giro {fmtv(g, 'x')} ≠ Receita/Capital investido = {fmtv(rol / c, 'x')}" + (' (estrangeira: pode ser câmbio)' if est else ''), rol / c)
    # lacunas no último ano
    for seg, s in D.items():
        last = max(y for y in s['metrics']['rol']['years'] if 'LTM' not in y)
        for e, ent in s['ents'].items():
            if e.startswith('dexco_'):
                continue
            tot = 0; falta = []
            for mk, m in s['metrics'].items():
                if last not in m['years'] or e not in m['series']:
                    continue
                if not [v for y, v in zip(m['years'], m['series'][e]) if y < last and v is not None]:
                    continue
                tot += 1
                if idx[(seg, mk, e, last)]['trat'] in ('sem dado', 'zero = sem dado'):
                    falta.append(m['label'])
            if falta:
                A.append(dict(pri='Alta' if e == 'dexco' else ('Média' if len(falta) >= tot / 2 else 'Baixa'), neg=SEG[seg], emp=ent['name'],
                              ind=f'{len(falta)} de {tot} indicadores', ano=last, v=None, fmt='', ref=None, alerta='Sem dado no último ano',
                              det='Sem dado (célula vazia, zero ou #N/D): ' + ', '.join(falta[:8]) + ('…' if len(falta) > 8 else ''), cel='', dex=e == 'dexco'))
    order = {'Alta': 0, 'Média': 1, 'Baixa': 2}
    A.sort(key=lambda a: (order[a['pri']], list(SEG.values()).index(a['neg']), not a['dex'], a['emp'], a['ind'], a['ano']))
    return A


def mudancas(R, R0):
    """Compara com a versão anterior: valores novos, alterados e que sumiram."""
    old = {(r['seg'], r['mk'], r['ent'], r['ano']): r for r in R0}
    out = []
    for r in R:
        k = (r['seg'], r['mk'], r['ent'], r['ano']); o = old.get(k)
        a, b = (o or {}).get('painel'), r['painel']
        if a is None and b is None:
            continue
        if a is None:
            tipo = 'Novo valor'
        elif b is None:
            tipo = 'Valor saiu do painel'
        elif abs(a - b) > 1e-6 * max(1, abs(a)):
            tipo = 'Valor alterado'
        else:
            continue
        out.append(dict(tipo=tipo, neg=SEG[r['seg']], emp=r['emp'], ind=r['ind'], ano=r['ano'], antes=a, depois=b, fmt=r['fmt'], cel=r['cel'],
                        trat=r['trat'] if b is None else ''))
    out.sort(key=lambda x: (['Valor alterado', 'Valor saiu do painel', 'Novo valor'].index(x['tipo']), x['neg'], x['emp'], x['ind'], x['ano']))
    return out


# ---------- 3. planilha ----------
NF = {'pct': '0.0%;-0.0%', 'x': '0.0"x"', 'd': '0" dias"', 'mi': '#,##0.0;(#,##0.0)'}
F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
HDR = PatternFill('solid', fgColor='111111'); YEL = PatternFill('solid', fgColor='FFF2A8'); DEX = PatternFill('solid', fgColor='ECEBE7')
LINE = Border(bottom=Side(style='thin', color='D9D8D3'))
STATUS = ['A conferir', 'OK', 'Corrigir na planilha', 'Investigar', 'Manter (exceção)']
TRAT = {'ok': 'ok', 'sem dado': 'sem dado', 'zero = sem dado': 'sem dado (célula zerada)', 'erro': 'retirado (erro) – n.s.',
        'outlier': 'retirado (outlier) – n.s.', 'n.s.': 'n.s. (EBITDA negativo)', 'calculado': 'calculado', 'alavancagem consolidada': 'consolidada Dexco'}


def header(ws, cols):
    for j, (h, w) in enumerate(cols, start=1):
        c = ws.cell(1, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = HDR
        c.alignment = Alignment(vertical='center', wrap_text=True); ws.column_dimensions[gl(j)].width = w


def raw_txt(x):
    if isinstance(x, (int, float)) or x is None:
        return x
    s = str(x)
    return 'erro na célula: ' + ERRTXT.get(s, 'erro') if s.startswith('#') else s


def planilha_saida(path, A, R, bad, nome_planilha, mud=None, em_aberto=None):
    wb = Workbook()
    ws = wb.active; ws.title = 'Leia-me'
    ws['A1'] = 'Conferência dos números – Análise de Peers'; ws['A1'].font = F(size=16, bold=True)
    ws['A2'] = f'Base: {nome_planilha}, abas "Gráficos" de cada negócio; valores exatamente como o painel os lê.'; ws['A2'].font = F(color='555555')
    txt = ['Como usar',
           '1. Comece pela aba "Pontos de atenção", ordenada por prioridade (Alta → Baixa). Use os filtros do cabeçalho para ver um negócio ou uma empresa.',
           '2. Para cada linha, abra a célula indicada em "Célula na planilha" e confira com a fonte (demonstração financeira / release).',
           '3. Preencha só as colunas amarelas: "Status" (lista) e "Comentário". O quadro abaixo se atualiza sozinho.',
           '4. "Base completa" traz todos os valores do painel com a célula de origem.' + (' "Mudanças" mostra o que mudou desde a versão anterior.' if mud is not None else ''),
           'Exemplo de linha preenchida: Status = "Corrigir na planilha" · Comentário = "Release 4T25: margem EBITDA ajustada é 18,2%; célula usa a reportada".',
           'Como os alertas foram gerados',
           '• Retirado do painel: erro de origem ou outlier pelas regras do painel (aparece como n.s. nas tabelas).',
           '• Salto em relação ao ano anterior: margens/ROIC/custos ≥ 15 p.p.; índices de caixa ≥ 100 p.p.; giro ≥ 0,8x; alavancagem ≥ 2,5x; prazos ≥ 60 dias; receita/capital ≥ 40%; EBITDA ≥ 60%.',
           '• Valor idêntico ao ano anterior e número redondo (múltiplo de R$ 100 mi): sinais de dado copiado ou estimado.',
           '• Sem dado no último ano: empresa com histórico, mas sem valor (vazio, zero ou #N/D) no ano mais recente.',
           '• Consistência entre blocos: margem EBITDA = EBITDA/receita, margem líquida = lucro/receita, giro = receita/capital investido.',
           f'• Verificação de leitura: {sum(1 for r in R if r["painel"] is not None)} valores do painel comparados com as células da planilha; {len(bad)} divergência(s).']
    r = 4
    for s in txt:
        tit = s in ('Como usar', 'Como os alertas foram gerados')
        c = ws.cell(r, 1, s); c.font = F(bold=tit, size=12 if tit else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
        if not tit:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); ws.row_dimensions[r].height = 28
        r += 1
    r += 1
    ws.cell(r, 1, 'Andamento da conferência').font = F(bold=True, size=12); r += 1
    for j, h in enumerate(['Negócio', 'Total', 'Alta'] + STATUS, start=1):
        c = ws.cell(r, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = HDR; c.alignment = Alignment(horizontal='center', wrap_text=True)
    r0 = r + 1; P = "'Pontos de atenção'"; n = len(A) + 1
    for i, neg in enumerate(SEG.values()):
        rr = r0 + i; ws.cell(rr, 1, neg).font = F()
        ws.cell(rr, 2, f'=COUNTIF({P}!$C$2:$C${n},$A{rr})')
        ws.cell(rr, 3, f'=COUNTIFS({P}!$C$2:$C${n},$A{rr},{P}!$B$2:$B${n},"Alta")')
        for j, s in enumerate(STATUS):
            ws.cell(rr, 4 + j, f'=COUNTIFS({P}!$C$2:$C${n},$A{rr},{P}!$L$2:$L${n},"{s}")')
    rt = r0 + 4; ws.cell(rt, 1, 'Total').font = F(bold=True)
    for j in range(2, 4 + len(STATUS)):
        ws.cell(rt, j, f'=SUM({gl(j)}{r0}:{gl(j)}{rt-1})')
        for rr in range(r0, rt + 1):
            ws.cell(rr, j).font = F(bold=rr == rt); ws.cell(rr, j).alignment = Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 16
    for j in range(2, 9):
        ws.column_dimensions[gl(j)].width = 13

    ws = wb.create_sheet('Pontos de atenção')
    header(ws, [('#', 5), ('Prioridade', 10), ('Negócio', 13), ('Empresa', 22), ('Indicador', 30), ('Ano', 8), ('Valor no painel', 15), ('Referência', 13),
                ('Alerta', 30), ('Detalhe', 60), ('Célula na planilha', 34), ('Status', 20), ('Comentário', 45)])
    for i, a in enumerate(A, start=2):
        for j, v in enumerate([i - 1, a['pri'], a['neg'], a['emp'], a['ind'], a['ano'], a['v'], a['ref'], a['alerta'], a['det'], a['cel'], 'A conferir', ''], start=1):
            c = ws.cell(i, j, v); c.font = F(); c.border = LINE; c.alignment = Alignment(vertical='top', wrap_text=j in (5, 10, 13))
        if a['fmt'] in NF:
            ws.cell(i, 7).number_format = ws.cell(i, 8).number_format = NF[a['fmt']]
        if a['dex']:
            for j in range(1, 12):
                ws.cell(i, j).fill = DEX
        ws.cell(i, 12).fill = YEL; ws.cell(i, 13).fill = YEL
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'L2:L{max(n, 2)}')
    ws.conditional_formatting.add(f'B2:B{max(n, 2)}', FormulaRule(formula=['$B2="Alta"'], font=Font(name='Arial', color='B3261E', bold=True)))
    ws.conditional_formatting.add(f'L2:L{max(n, 2)}', FormulaRule(formula=['$L2="OK"'], fill=PatternFill('solid', fgColor='D7F0DD')))
    ws.conditional_formatting.add(f'L2:L{max(n, 2)}', FormulaRule(formula=['$L2="Corrigir na planilha"'], fill=PatternFill('solid', fgColor='F8D3CF')))
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:M{max(n, 2)}'
    ws.cell(n + 2, 1, 'Linhas em cinza = Dexco. "Referência" = valor do ano anterior (saltos) ou valor recalculado (consistência).').font = F(italic=True, color='555555')

    if mud is not None:
        ws = wb.create_sheet('Mudanças')
        header(ws, [('Tipo', 20), ('Negócio', 13), ('Empresa', 24), ('Indicador', 32), ('Ano', 9), ('Antes', 15), ('Depois', 15), ('Célula na planilha', 36), ('Observação', 30)])
        for i, x in enumerate(mud, start=2):
            for j, v in enumerate([x['tipo'], x['neg'], x['emp'], x['ind'], x['ano'], x['antes'], x['depois'], x['cel'], TRAT.get(x['trat'], x['trat'])], start=1):
                ws.cell(i, j, v).font = F()
            if x['fmt'] in NF:
                ws.cell(i, 6).number_format = ws.cell(i, 7).number_format = NF[x['fmt']]
        ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:I{len(mud) + 1}'
        if not mud:
            ws.cell(2, 1, 'Nenhuma mudança em relação à versão anterior.').font = F(italic=True)

    ws = wb.create_sheet('Base completa')
    header(ws, [('Negócio', 13), ('Empresa', 26), ('Indicador', 34), ('Ano', 9), ('Valor no painel', 15), ('Valor na célula', 15), ('Célula na planilha', 36), ('Tratamento no painel', 24), ('Motivo', 55)])
    for i, x in enumerate(R, start=2):
        for j, v in enumerate([SEG[x['seg']], x['emp'], x['ind'], x['ano'], x['painel'], raw_txt(x['planilha']), x['cel'], TRAT.get(x['trat'], x['trat']), x.get('why', '')], start=1):
            ws.cell(i, j, v).font = F()
        if x['fmt'] in NF:
            ws.cell(i, 5).number_format = ws.cell(i, 6).number_format = NF[x['fmt']]
    ws.freeze_panes = 'E2'; ws.auto_filter.ref = f'A1:I{len(R) + 1}'

    if em_aberto:
        ws = wb.create_sheet('Em aberto')
        header(ws, [('Tema', 26), ('O que vimos', 70), ('Decisão até agora', 40), ('Status', 18), ('Comentário', 45)])
        for i, row in enumerate(em_aberto, start=2):
            for j, v in enumerate([row.get('tema'), row.get('visto'), row.get('decisao'), 'A conferir', ''], start=1):
                c = ws.cell(i, j, v); c.font = F(bold=j == 1); c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.cell(i, 4).fill = YEL; ws.cell(i, 5).fill = YEL
        dv2 = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'D2:D{len(em_aberto) + 1}')
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--projeto', default='.', help='pasta do projeto painel-estrategico (tem pipeline/peers)')
    ap.add_argument('--planilha', required=True, help='Análise de Competidores (versão a conferir)')
    ap.add_argument('--anterior', help='versão anterior da planilha, para a aba "Mudanças"')
    ap.add_argument('--em-aberto', help='JSON com pontos em aberto: [{"tema","visto","decisao"}]')
    ap.add_argument('--saida', default='Conferencia_Peers.xlsx')
    a = ap.parse_args()
    D = pipeline(a.projeto, a.planilha)
    R, bad = linhas(D, a.planilha)
    A = alertas(D, R)
    mud = None
    if a.anterior:
        D0 = pipeline(a.projeto, a.anterior); R0, _ = linhas(D0, a.anterior); mud = mudancas(R, R0)
    ab = json.load(open(a.em_aberto, encoding='utf-8')) if a.em_aberto else None
    planilha_saida(a.saida, A, R, bad, os.path.basename(a.planilha), mud, ab)
    resumo = {'saida': os.path.abspath(a.saida), 'valores_conferidos': sum(1 for r in R if r['painel'] is not None),
              'divergencias_de_leitura': len(bad), 'alertas': len(A),
              'por_prioridade': dict(collections.Counter(x['pri'] for x in A)),
              'por_tipo': dict(collections.Counter(x['alerta'] for x in A)),
              'alta_dexco': [f"{x['neg']} · {x['emp']} · {x['ind']} {x['ano']}: {x['det']}" for x in A if x['pri'] == 'Alta'][:30]}
    if mud is not None:
        resumo['mudancas'] = dict(collections.Counter(x['tipo'] for x in mud))
    if bad:
        resumo['exemplos_divergencia'] = bad[:10]
    print(json.dumps(resumo, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
